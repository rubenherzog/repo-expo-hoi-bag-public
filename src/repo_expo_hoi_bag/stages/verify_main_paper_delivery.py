#!/usr/bin/env python3
"""Verify the lightweight local delivery for the main k10 paper."""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook


CHECKOUT_ROOT = Path(__file__).resolve().parents[3]
CONFIG = CHECKOUT_ROOT / "config" / "main_paper_delivery.yaml"


def _args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=CONFIG)
    parser.add_argument("--write-manifest", action="store_true")
    parser.add_argument(
        "--refresh-manifest",
        action="store_true",
        help="Explicitly replace only the generated status manifest.",
    )
    return parser.parse_args()


def _figure_status(root: Path, stem: str) -> tuple[bool, str]:
    matches = list(root.rglob(f"{stem}.*"))
    extensions = {path.suffix.lower() for path in matches}
    missing = {".pdf", ".png", ".svg"}.difference(extensions)
    source_data = list(root.rglob(f"{stem}_source_data.xlsx"))
    if missing or not source_data:
        detail = []
        if missing:
            detail.append(f"formats={','.join(sorted(missing))}")
        if not source_data:
            detail.append("source_data_xlsx")
        return False, "; ".join(detail)
    return True, "complete"


def _validate_whole_pca_source_data(root: Path) -> None:
    """Reject the stale BAG-specific PCA basis used before effective-sample dedup."""
    source_root = root / "figures" / "supplementary" / "source_data"
    stems = {
        "structural": "whole_exposome_pca_sensitivity_source_data_a1_struct_pca_variance.csv",
        "functional": "whole_exposome_pca_sensitivity_source_data_b1_func_pca_variance.csv",
    }
    frames = {bag: pd.read_csv(source_root / name) for bag, name in stems.items()}
    columns = [
        "pc_n",
        "explained_variance_ratio",
        "cumulative_explained_variance_ratio",
    ]
    for bag, frame in frames.items():
        missing = set(columns).difference(frame.columns)
        if missing:
            raise ValueError(f"Whole-exposome PCA source for {bag} is missing {sorted(missing)}")
    structural = frames["structural"][columns].reset_index(drop=True)
    functional = frames["functional"][columns].reset_index(drop=True)
    if not structural["pc_n"].equals(functional["pc_n"]) or not np.allclose(
        structural[columns[1:]].to_numpy(float),
        functional[columns[1:]].to_numpy(float),
        rtol=0.0,
        atol=1e-12,
    ):
        raise ValueError(
            "Whole-exposome PCA variance differs by BAG; both rows must use the same "
            "effective country-year exposome matrix"
        )


def main() -> None:
    args = _args()
    config = yaml.safe_load(args.config.read_text(encoding="utf-8"))
    root = CHECKOUT_ROOT / config["delivery_root"]
    records: list[dict[str, str]] = []
    for group, rel in (("main_figure", "figures/main"), ("supplementary_figure", "figures/supplementary")):
        for stem in config[f"{group}s"]:
            complete, detail = _figure_status(root / rel, stem)
            records.append({"kind": group, "name": stem, "status": "complete" if complete else "missing", "detail": detail})
    workbook = root / "tables" / config["supplementary_tables"]["workbook"]
    if workbook.is_file():
        sheets = load_workbook(workbook, read_only=True).sheetnames
        expected = config["supplementary_tables"]["sheets"]
        missing_sheets = [sheet for sheet in expected if sheet not in sheets]
        status = "complete" if not missing_sheets else "incomplete"
        detail = "complete" if not missing_sheets else f"missing_sheets={','.join(missing_sheets)}"
    else:
        status, detail = "missing", "workbook"
    records.append({"kind": "supplementary_tables", "name": config["supplementary_tables"]["workbook"], "status": status, "detail": detail})
    _validate_whole_pca_source_data(root)
    manifest = root / "MANIFEST.csv"
    if args.write_manifest or args.refresh_manifest:
        if manifest.exists() and not args.refresh_manifest:
            raise FileExistsError(f"Refusing to overwrite delivery manifest: {manifest}")
        root.mkdir(parents=True, exist_ok=True)
        with manifest.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=["kind", "name", "status", "detail"])
            writer.writeheader()
            writer.writerows(records)
    missing = [record for record in records if record["status"] != "complete"]
    for record in records:
        print(f"{record['status']:10} {record['kind']:22} {record['name']} {record['detail']}")
    if missing:
        raise SystemExit(f"Main-paper delivery incomplete: {len(missing)} required entries missing or incomplete")
    print(f"Main-paper delivery verified: {root}")


if __name__ == "__main__":
    main()
