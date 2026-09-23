"""Append the education-only sensitivity summary to delivered ST17.

The no-scanner figure is a delivery-specific extension of the education/scanner
sensitivity.  This updater keeps the original ST17 sections intact and adds
its sample accounting plus the country-balanced best-single results used by
the orange marks in the new figure.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
RUNTIME = Path("/data/workspaces/neuromodelling/rherzog/Brainlat/expo_hoi_bag_repro_data")
FIG_SOURCE = ROOT / "outputs/main/paper/complete/figures/supplementary/source_data"
TABLE_ROOT = ROOT / "outputs/main/paper/complete/tables"
SOURCE_DIR = TABLE_ROOT / "source_data/selection_sensitivities"
WORKBOOK = TABLE_ROOT / "Supplementary_Tables_main_k10.xlsx"
EDU_EVAL = RUNTIME / "work/analysis_runs/main_k10_release_20260916/sensitivity_eval/education_scanner_baseline"
FULL_ROOT = RUNTIME / "results/analysis_runs/paper_reanalysis_k10"

NAVY = "1F4E78"
BLUE = "D9EAF7"
WHITE = "FFFFFF"
THIN_GREY = Side(style="thin", color="B7B7B7")


def _sample_frame() -> pd.DataFrame:
    rows = []
    for bag, label in (("structural", "Structural"), ("functional", "Functional")):
        education = pd.read_csv(EDU_EVAL / bag / "plus_education" / f"{bag}_ols_country.csv")
        education = education[education["candidate_id"].astype(str).eq("__baseline__")]
        full = pd.read_csv(FULL_ROOT / "ols" / bag / "ols" / "baseline" / "metrics_country.csv")
        education_n = int(education["n_test"].sum())
        full_n = int(full["n_test"].sum())
        rows.append({
            "BAG": label,
            "Education analysis participants": education_n,
            "Education analysis countries": int(education["fold_country"].nunique()),
            "Full analysis participants": full_n,
            "Full analysis countries": int(full["fold_country"].nunique()),
            "Participants retained (%)": 100.0 * education_n / full_n,
            "Participants excluded": full_n - education_n,
        })
    return pd.DataFrame(rows)


def _best_single_frame() -> pd.DataFrame:
    rows = []
    for bag, label, panel in (
        ("structural", "Structural", "a1_struct_edu"),
        ("functional", "Functional", "b1_func_edu"),
    ):
        frame = pd.read_csv(FIG_SOURCE / f"education_scanner_baseline_no_scanner_source_data_{panel}.csv")
        best = frame[["rung_id", "best_single_candidate_id", "predictors_identity", "best_single_r2"]].drop_duplicates()
        for row in best.itertuples(index=False):
            rows.append({
                "BAG": label,
                "Model level": {"ols": "OLS", "xgb_tree_d1": "d1", "xgb_tree_d2": "d2", "xgb_tree_d3": "d3"}[row.rung_id],
                "Best single exposure": row.best_single_candidate_id,
                "Exposure": row.predictors_identity,
                "LOCO R² (unweighted country mean)": float(row.best_single_r2),
                "Participants": 7331 if bag == "structural" else 6206,
                "Countries": 20 if bag == "structural" else 13,
                "Covariates": "Baseline covariates + Education",
                "XGBoost HPO scope": "Not applicable (OLS)" if row.rung_id == "ols" else "single_exposure",
            })
    return pd.DataFrame(rows)


def _append_section(ws, row: int, title: str, frame: pd.DataFrame) -> int:
    ws.cell(row, 1, title).font = Font(bold=True, size=11)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=BLUE)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(frame.columns))
    row += 1
    for col, name in enumerate(frame.columns, 1):
        cell = ws.cell(row, col, name)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=THIN_GREY)
    row += 1
    for values in frame.itertuples(index=False, name=None):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
            if isinstance(value, float):
                cell.number_format = "0.000000"
        row += 1
    ws.auto_filter.ref = f"A{row - len(frame) - 1}:{get_column_letter(len(frame.columns))}{row - 1}"
    return row + 1


def main() -> None:
    if not WORKBOOK.is_file():
        raise FileNotFoundError(WORKBOOK)
    samples = _sample_frame()
    singles = _best_single_frame()
    SOURCE_DIR.mkdir(parents=True, exist_ok=True)
    source = SOURCE_DIR / "education_no_scanner_summary.csv"
    pd.concat([
        samples.assign(section="Sample size comparison"),
        singles.assign(section="Best single + Education"),
    ], ignore_index=True, sort=False).to_csv(source, index=False)

    book = load_workbook(WORKBOOK)
    ws = book["ST17_CovariatesTargets"]
    # Remove a prior run of this exact addendum without touching the established
    # covariate or residualised-target sections above it.
    marker = "Education-only no-scanner figure"
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == marker:
            ws.delete_rows(row, ws.max_row - row + 1)
            break
    row = ws.max_row + 2
    row = _append_section(ws, row, marker, samples)
    _append_section(ws, row, "Best single exposure with Education", singles)
    book.save(WORKBOOK)
    print(f"Updated {WORKBOOK} and wrote {source}")


if __name__ == "__main__":
    main()
