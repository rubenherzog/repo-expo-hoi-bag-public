"""Build 18 compact supplementary tables for transfer to the paper.

The canonical ``Supplementary_Tables.xlsx`` remains the exhaustive audit
workbook.  This stage creates a separate presentation workbook with exactly
one rectangular, portrait-oriented table per worksheet and exactly the
Supplementary Table 1--18 numbering used by the manuscript.  It performs no
statistical fitting and reads values only from existing reporting artefacts.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from statsmodels.stats.multitest import multipletests


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "docs" / "supplementary" / "Supplementary_Tables.xlsx"
OUTPUT = ROOT / "docs" / "supplementary" / "Supplementary_Tables_Publication.xlsx"

FONT = "Times New Roman"
BLACK = "000000"
GREY = "666666"
RULE = Side(style="thin", color=BLACK)


@dataclass(frozen=True)
class PublicationTable:
    number: int
    title: str
    frame: pd.DataFrame
    note: str

    @property
    def sheet_name(self) -> str:
        return f"ST{self.number:02d}"

    @property
    def label(self) -> str:
        return f"Supplementary Table {self.number}"


def _section_frame(worksheet, heading: str) -> pd.DataFrame:
    heading_row = next(
        (row for row in range(1, worksheet.max_row + 1) if worksheet.cell(row, 1).value == heading),
        None,
    )
    if heading_row is None:
        raise ValueError(f"Section {heading!r} not found in {worksheet.title}")
    header_row = heading_row + 1
    headers = [worksheet.cell(header_row, column).value for column in range(1, worksheet.max_column + 1)]
    width = max(index for index, value in enumerate(headers, start=1) if value is not None)
    rows: list[list[object]] = []
    for row in range(header_row + 1, worksheet.max_row + 1):
        values = [worksheet.cell(row, column).value for column in range(1, width + 1)]
        if all(value is None for value in values):
            break
        rows.append(values)
    return pd.DataFrame(rows, columns=headers[:width])


def _normalise(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.rename(
        columns={
            "Bag": "BAG",
            "Objective": "Discovery arm",
            "N Subjects": "Participants",
            "Bias Mean": "Mean bias (years)",
            "Mae": "MAE (years)",
            "Rmse": "RMSE (years)",
            "P Value": "P value",
        }
    ).copy()
    if "BAG" in frame:
        frame["BAG"] = frame["BAG"].replace(
            {"structural": "Structural", "functional": "Functional"}
        )
        order = frame["BAG"].map({"Structural": 0, "Functional": 1}).fillna(2)
        frame = frame.assign(_bag_order=order).sort_values("_bag_order", kind="stable").drop(columns="_bag_order")
    for column in ("Discovery arm", "High-order arm"):
        if column in frame:
            frame[column] = frame[column].replace(
                {"o_min": "Synergy arm", "o_max": "Redundancy arm", "syn": "Synergy arm", "red": "Redundancy arm"}
            )
    if "Diagnosis" in frame:
        frame["Diagnosis"] = frame["Diagnosis"].replace({"CN": "HC", "FTD": "FTLD"})
    return frame.reset_index(drop=True)


def _get(source, sheet: str, section: str) -> pd.DataFrame:
    return _normalise(_section_frame(source[sheet], section))


def _ci(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    if {"95% CI, lower", "95% CI, upper"}.issubset(frame.columns):
        position = frame.columns.get_loc("95% CI, lower")
        values = frame.apply(
            lambda row: f"{row['95% CI, lower']:.3f} to {row['95% CI, upper']:.3f}", axis=1
        )
        frame = frame.drop(columns=["95% CI, lower", "95% CI, upper"])
        frame.insert(position, "95% CI", values)
    return frame


def _concat_records(frames: list[pd.DataFrame], columns: list[str]) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    for frame in frames:
        records.extend(frame.reindex(columns=columns).to_dict(orient="records"))
    return pd.DataFrame.from_records(records, columns=columns)


def _holm_adjust(
    frame: pd.DataFrame,
    p_column: str,
    group_columns: list[str],
    *,
    adjusted_column: str,
) -> pd.DataFrame:
    """Apply Holm correction within explicitly defined manuscript families."""
    result = frame.copy()
    result[adjusted_column] = None
    valid = result[p_column].notna()
    groups = (
        result.loc[valid].groupby(group_columns, sort=False, dropna=False).groups
        if group_columns
        else {"all": result.index[valid]}
    )
    for indices in groups.values():
        result.loc[indices, adjusted_column] = multipletests(
            result.loc[indices, p_column].astype(float), method="holm"
        )[1]
    return result


def _table_01(source) -> PublicationTable:
    regression = _ci(_get(source, "ST01_HigherOrder", "Sub-combination enrichment regressions"))
    baseline = _get(source, "ST01_HigherOrder", "Exposome-wide negative-Ω baselines")
    baseline["Exposome-wide negative-Ω baseline"] = baseline.apply(
        lambda row: (
            f"{int(row['Negative-Ω combinations']):,}/{int(row['Total combinations']):,} "
            f"({row['Negative-Ω fraction']:.6f})"
        ),
        axis=1,
    )
    baseline = baseline[["Combination order", "Exposome-wide negative-Ω baseline"]]
    frame = regression.merge(baseline, on="Combination order", how="left")
    frame["Term"] = frame["Term"].replace(
        {"is_syn": "Synergy-arm coefficient", "thoi_o": "Aggregate O-information"}
    )
    frame = _holm_adjust(
        frame,
        "Permutation p",
        ["BAG"],
        adjusted_column="Holm-adjusted permutation p",
    )
    frame = frame[["BAG", "Combination order", "Term", "Estimate", "SE", "95% CI", "Permutation p", "Holm-adjusted permutation p", "Candidates", "Exposome-wide negative-Ω baseline"]]
    return PublicationTable(1, "Higher-order information and sub-combination enrichment", frame, "Negative-O-information fractions are exposome-wide baselines. Arm coefficients were tested with 10,000 label permutations within model level; P values were Holm-adjusted across combination orders within each BAG measure. Parametric candidate-level P values were not used because candidate sets overlap. Full set-size trajectories are provided in the Source Data for Fig. 2.")


def _table_02(source) -> PublicationTable:
    multi = _get(source, "ST02_ModelComplexity", "Best multivariate candidate at each model level")
    multi = multi.pivot(index=["BAG", "Model level"], columns="Discovery arm", values=["Best LOCO R²", "Selected set size"])
    multi.columns = [f"{arm.replace(' arm', '')} {metric.lower()}" for metric, arm in multi.columns]
    multi = multi.reset_index()
    single = _get(source, "ST02_ModelComplexity", "Best single exposure at each model level")
    single = single[["BAG", "Model level", "Selected exposure", "Best single-exposure LOCO R²"]]
    frame = multi.merge(single, on=["BAG", "Model level"], validate="one_to_one")
    test = _ci(_get(source, "ST02_ModelComplexity", "Paired d3-versus-OLS single-exposure comparison"))
    test = test[["BAG", "ΔR², d3 − OLS", "95% CI", "Country-cluster bootstrap p"]]
    test = _holm_adjust(
        test,
        "Country-cluster bootstrap p",
        [],
        adjusted_column="Holm-adjusted country-cluster bootstrap p",
    )
    frame = frame.merge(test, on="BAG", how="left")
    frame.loc[frame["Model level"] != "d3", ["ΔR², d3 − OLS", "95% CI", "Country-cluster bootstrap p", "Holm-adjusted country-cluster bootstrap p"]] = None
    for arm in ("Redundancy", "Synergy"):
        frame[f"{arm} R² (set size)"] = frame.apply(
            lambda row: f"{row[f'{arm} best loco r²']:.3f} ({int(row[f'{arm} selected set size'])})",
            axis=1,
        )
    frame = frame[
        [
            "BAG",
            "Model level",
            "Redundancy R² (set size)",
            "Synergy R² (set size)",
            "Selected exposure",
            "Best single-exposure LOCO R²",
            "ΔR², d3 − OLS",
            "95% CI",
            "Country-cluster bootstrap p",
            "Holm-adjusted country-cluster bootstrap p",
        ]
    ]
    level_order = {"OLS": 0, "d1": 1, "d2": 2, "d3": 3}
    frame = frame.assign(_level=frame["Model level"].map(level_order)).sort_values(["BAG", "_level"], key=lambda col: col.map({"Structural": 0, "Functional": 1}) if col.name == "BAG" else col).drop(columns="_level")
    return PublicationTable(2, "Predictive performance across model-complexity levels", frame.reset_index(drop=True), "Multivariate candidates and single exposures were selected independently within each model level by pooled LOCO R². The d3-versus-OLS single-exposure comparison used paired out-of-fold predictions and 10,000 country-cluster bootstrap samples; P values were Holm-adjusted across the two BAG measures. The comparison does not isolate complexity when the selected exposure differs.")


def _table_03(source) -> PublicationTable:
    tests = _ci(_get(source, "ST03_BestModels", "Matched depth-3 R² comparisons"))
    effects = _get(source, "ST03_BestModels", "Incremental effect sizes versus covariate baseline")
    frame = tests.merge(effects, on=["BAG", "Comparison (A vs B)"], how="left")
    frame["_family"] = frame["Comparison (A vs B)"].map(
        lambda value: "baseline" if "covariate baseline" in value else "architecture"
    )
    frame = _holm_adjust(
        frame,
        "Country-cluster bootstrap p",
        ["BAG", "_family"],
        adjusted_column="Holm-adjusted country-cluster bootstrap p",
    ).drop(columns="_family")

    complexity = pd.read_csv(
        ROOT / "outputs" / "dedup" / "model_comparison" / "complexity_pairwise_comparisons.csv"
    )
    complexity["BAG"] = complexity["bag"].str.title()
    complexity["Comparison (A vs B)"] = complexity.apply(
        lambda row: f"{row['trajectory']}: {row['comparison_a_vs_b']}",
        axis=1,
    )
    complexity_rows = pd.DataFrame(
        {
            "BAG": complexity["BAG"],
            "Comparison (A vs B)": complexity["Comparison (A vs B)"],
            "Participants": complexity["n_subjects"],
            "Countries": complexity["n_countries"],
            "R², model A": complexity["r2_a"],
            "R², model B": complexity["r2_b"],
            "ΔR², A − B": complexity["delta_r2_a_minus_b"],
            "95% CI": complexity.apply(
                lambda row: (
                    f"{row['delta_r2_ci_lower']:.3f} to "
                    f"{row['delta_r2_ci_upper']:.3f}"
                ),
                axis=1,
            ),
            "Country-cluster bootstrap p": complexity[
                "country_cluster_bootstrap_p_raw"
            ],
            "Holm-adjusted country-cluster bootstrap p": complexity[
                "holm_p_within_bag_trajectory"
            ],
            "Cohen f²": None,
        }
    )
    frame = _concat_records([frame, complexity_rows], list(frame.columns))
    frame = frame[
        [
            "BAG",
            "Comparison (A vs B)",
            "Participants",
            "Countries",
            "R², model A",
            "R², model B",
            "ΔR², A − B",
            "95% CI",
            "Country-cluster bootstrap p",
            "Holm-adjusted country-cluster bootstrap p",
            "Cohen f²",
        ]
    ]
    return PublicationTable(3, "Matched comparisons of deployed models", frame, "Comparisons used matched out-of-fold predictions and 10,000 country-cluster bootstrap samples. For deployed depth-3 comparisons, P values were Holm-adjusted within baseline or architecture families for each BAG measure. Model-level P values were Holm-adjusted across the six pairwise level comparisons within each BAG measure and trajectory. Level-selected rows compare independently selected winners; fixed-d3-winner rows hold the deployed d3 candidate set constant across levels. Cohen's f² is reported only for comparisons with the covariate baseline.")


def _table_04(source) -> PublicationTable:
    null = _get(source, "ST04_CountryBlockNull", "Depth-3 deployed models")
    cross = _get(source, "ST04_CountryBlockNull", "Country-cross-fitted candidate-selection sensitivity")
    null["Discovery arm"] = null["Model"].str.extract(r"(synergy|redundancy)", expand=False).str.title() + " arm"
    cross["Positive held-out folds"] = cross["Countries with R² > 0"].astype(int).astype(str) + "/" + cross["Held-out countries"].astype(int).astype(str)
    frame = null.merge(cross, on=["BAG", "Discovery arm"], how="left")
    frame = _holm_adjust(
        frame,
        "Empirical upper-tail p",
        ["BAG"],
        adjusted_column="Holm-adjusted upper-tail p",
    )
    frame = frame[["BAG", "Model", "Set size", "Observed R²", "Null mean R²", "z versus null", "Empirical upper-tail p", "Holm-adjusted upper-tail p", "Most frequent candidate fraction", "Median held-out-country R²", "Positive held-out folds"]]
    return PublicationTable(4, "Country-block permutation null and selection stability", frame, "Country-year exposure signatures were permuted as blocks while participants, BAG, covariates and LOCO folds were fixed (10,000 permutations). P values were Holm-adjusted across the three deployed models within each BAG measure. Cross-fitted summaries select each candidate without using its evaluation country; fold-level records remain in the audit workbook.")


def _table_05(source) -> PublicationTable:
    frame = _get(source, "ST05_TopComposition", "Set size and synergistic-triplet content")
    frontier = _get(source, "ST05_TopComposition", "Top-frontier arm enrichment")
    frontier = frontier.rename(
        columns={
            "Models": "Top-frontier models",
            "Synergy-arm fraction": "Top-frontier synergy fraction",
        }
    )
    frame = frame.merge(
        frontier[
            [
                "BAG",
                "Top-frontier models",
                "Top-frontier synergy fraction",
                "Enrichment ratio",
            ]
        ],
        on="BAG",
        validate="many_to_one",
    )
    frame["Exposome-wide negative-Ω triplet fraction"] = frame["Metric"].map(
        lambda metric: 0.332301 if metric == "Negative-Ω triplet fraction" else None
    )
    frame = _holm_adjust(
        frame,
        "Two-sided label-permutation p",
        ["BAG"],
        adjusted_column="Holm-adjusted permutation p",
    )
    return PublicationTable(5, "Composition of top-performing candidate sets", frame, "Medians summarize the top 50 depth-3 candidates within each BAG measure and arm; two-sided P values use 10,000 overlap-aware arm-label permutations and were Holm-adjusted across the two composition metrics within each BAG measure. Performance-frontier columns summarize the top 10% of candidates and are descriptive; the candidate-pool synergy fraction was 0.5. Candidate-level rows remain in the audit workbook.")


def _table_06(source) -> PublicationTable:
    ranges = _get(source, "ST06_DomainDiversity", "Observed candidate-level performance ranges")
    correlations = _get(source, "ST06_DomainDiversity", "Arm-specific Pearson correlations")
    interaction = _get(source, "ST06_DomainDiversity", "Interaction permutation test")
    frame = ranges.merge(correlations, on=["BAG", "Discovery arm", "Candidates"], validate="one_to_one")
    slopes = []
    for row in frame.itertuples(index=False):
        bag_result = interaction[interaction["BAG"] == row[0]].iloc[0]
        arm = row[1]
        slopes.append(bag_result["Synergy-arm slope" if arm == "Synergy arm" else "Redundancy-arm slope"])
    frame["Adjusted diversity slope"] = slopes
    frame = frame.merge(interaction[["BAG", "Synergy − redundancy interaction", "Within-order permutation p"]], on="BAG", how="left")
    unique_p = interaction.drop_duplicates("BAG")[["BAG", "Within-order permutation p"]]
    unique_p = _holm_adjust(
        unique_p,
        "Within-order permutation p",
        [],
        adjusted_column="Holm-adjusted permutation p",
    ).set_index("BAG")["Holm-adjusted permutation p"]
    frame["Holm-adjusted permutation p"] = frame["BAG"].map(unique_p)
    return PublicationTable(6, "Domain diversity and predictive performance", frame, "Pearson correlations and unadjusted slopes are descriptive. Adjusted slopes control for set size; arm interactions were tested by 10,000 arm-label permutations within set size and Holm-adjusted across the two BAG measures.")


def _table_07(source) -> PublicationTable:
    frame = _get(source, "ST07_DomainComposition", "Domain counts and percentages")
    return PublicationTable(7, "Domain composition of top-performing models", frame, "Counts describe the top 20 depth-3 candidates in each BAG measure and discovery arm. Percentages use all feature occurrences or all 20 candidate sets, as indicated.")


def _table_08(source) -> PublicationTable:
    frame = _get(source, "ST08_Cooccurrence", "Observed network contrasts")
    frame["Network statistic"] = frame["Network statistic"].replace(
        {
            "n_domains": "Distinct domains",
            "n_edges": "Domain-pair edges",
            "density": "Network density",
            "mean_weighted_degree": "Mean weighted degree",
            "prevalence_gini": "Domain-prevalence Gini coefficient",
            "edge_weight_gini": "Edge-weight Gini coefficient",
        }
    )
    frame = _holm_adjust(
        frame,
        "Two-sided permutation p",
        ["BAG"],
        adjusted_column="Holm-adjusted permutation p",
    )
    return PublicationTable(8, "Domain co-occurrence network architecture", frame, "Networks were constructed from the top 20 depth-3 candidates per BAG measure and arm. P values use 10,000 overlap-aware arm-label permutations and were Holm-adjusted across the six network statistics within each BAG measure.")


def _table_09(source) -> PublicationTable:
    top = _get(source, "ST09_RecurrentTriplets", "Most recurrent triplet by BAG and arm")
    top["Category"] = "Most recurrent"
    top["Best single exposure"] = None
    top["Single-exposure R²"] = None
    single = _get(source, "ST09_RecurrentTriplets", "Triplets containing the best single exposure")
    single["Category"] = "Contains best single exposure"
    common = ["BAG", "Discovery arm", "Category", "Best single exposure", "Single-exposure R²", "Triplet", "Candidates containing triplet", "Candidates in arm", "Prevalence (%)", "Ω (nats)"]
    frame = _concat_records([top, single], common)
    shared = _get(source, "ST09_RecurrentTriplets", "Triplets shared across both BAGs and arms")
    highlighted = shared[
        shared["Triplet"].eq("Nitrogen oxide (NOx)|Sulphur dioxide (SO₂) emissions|soc_grp_equal_est")
    ].iloc[0]
    shared_row = {
        "BAG": "Both",
        "Discovery arm": "Both",
        "Category": "Highlighted triplet shared across all groups",
        "Best single exposure": None,
        "Single-exposure R²": None,
        "Triplet": highlighted["Triplet"],
        "Candidates containing triplet": None,
        "Candidates in arm": None,
        "Prevalence (%)": (
            f"Structural: synergy {highlighted['Synergy arm, structural BAG (%)']:.3f}, "
            f"redundancy {highlighted['Redundancy arm, structural BAG (%)']:.3f}; "
            f"Functional: synergy {highlighted['Synergy arm, functional BAG (%)']:.3f}, "
            f"redundancy {highlighted['Redundancy arm, functional BAG (%)']:.3f}"
        ),
        "Ω (nats)": highlighted["Ω (nats)"],
    }
    frame = pd.DataFrame.from_records([*frame.to_dict(orient="records"), shared_row], columns=common)
    return PublicationTable(9, "Recurrent exposome triplets", frame, "Triplets were counted across the top 20 candidates at each model level, counting a candidate once if it recurred across levels. Eleven triplets occurred in both arms for both BAG measures; one representative is highlighted, and the exhaustive listing remains in the audit workbook.")


def _table_10(source) -> PublicationTable:
    frame = _ci(_get(source, "ST10_ResidualBias", "Primary within-country diagnosis contrasts"))
    frame = _holm_adjust(
        frame,
        "Two-sided sign-flip p",
        ["BAG", "Discovery arm"],
        adjusted_column="Holm-adjusted sign-flip p",
    )
    frame = frame[["BAG", "Discovery arm", "Outcome", "Diagnosis contrast", "Matched countries", "Mean country difference", "95% CI", "Two-sided sign-flip p", "Holm-adjusted sign-flip p"]]
    return PublicationTable(10, "Country-aware diagnostic contrasts in BAG residual bias", frame, "Bias is predicted minus observed BAG, in years. Disease-minus-HC differences were calculated within countries containing both groups and averaged with equal country weights; intervals used 10,000 country-bootstrap samples and P values used 10,000 sign-flip permutations. P values were Holm-adjusted across the two diagnoses and two outcomes within each BAG measure and arm. Participant-level summaries are provided in Fig. 4 Source Data.")


def _table_11(source) -> PublicationTable:
    frame = _get(source, "ST11_NormativeTransfer", "Diagnostic train-to-test performance")
    frame = _holm_adjust(
        frame,
        "Wilcoxon p",
        ["BAG"],
        adjusted_column="Holm-adjusted Wilcoxon p",
    )
    frame = _holm_adjust(
        frame,
        "Sign-test p",
        ["BAG"],
        adjusted_column="Holm-adjusted sign-test p",
    )
    frame["Holm-adjusted p (Wilcoxon; sign)"] = frame.apply(
        lambda row: (
            f"{row['Holm-adjusted Wilcoxon p']:.6g}; "
            f"{row['Holm-adjusted sign-test p']:.6g}"
        ),
        axis=1,
    )
    frame = frame.drop(
        columns=["Holm-adjusted Wilcoxon p", "Holm-adjusted sign-test p"]
    )
    frame = frame[["BAG", "Train → test", "Participants", "Countries", "Baseline R²", "Best arm", "Best R²", "ΔR², best − baseline", "Wilcoxon p", "Sign-test p", "Holm-adjusted p (Wilcoxon; sign)"]]
    return PublicationTable(11, "Diagnostic train-to-test performance", frame, "Depth-3 models were selected independently within each train-to-test setting. Across-country tests were paired and two-sided; P values for each test were Holm-adjusted across the seven train-to-test settings within each BAG measure.")


def _table_12(source) -> PublicationTable:
    fixed = _ci(_get(source, "ST12_MixedModels", "Fixed effects"))
    fixed["Analysis"] = "Fixed effect"
    fixed["Term"] = fixed["Fixed effect"]
    fixed["P value"] = fixed["Two-sided Wald p"]
    fixed["Sample"] = fixed["Participants"].astype(int).map("{:,}".format) + " participants; " + fixed["Countries"].astype(int).astype(str) + " countries"
    fixed["Additional statistic"] = None
    variance = _get(source, "ST12_MixedModels", "Variance components and country random-intercept tests")
    variance["Analysis"] = "Random intercept"
    variance["Term"] = "Country"
    variance["Estimate"] = variance["Adjusted country variance"]
    variance["SE"] = None
    variance["95% CI"] = None
    variance["P value"] = variance["Parametric-bootstrap p"]
    variance["Sample"] = variance["Participants"].astype(int).map("{:,}".format) + " participants; " + variance["Countries"].astype(int).astype(str) + " countries"
    variance["Additional statistic"] = variance.apply(
        lambda row: f"Adjusted ICC={row['Adjusted ICC']:.3f}; LRT={row['Random-intercept LRT']:.3f}",
        axis=1,
    )
    columns = ["BAG", "Discovery arm", "Analysis", "Term", "Estimate", "SE", "95% CI", "P value", "Sample", "Additional statistic"]
    frame = _concat_records([fixed, variance], columns)
    frame["Holm-adjusted p"] = None
    diagnosis_mask = frame["Term"].isin(["diagnosis__AD", "diagnosis__FTLD"])
    winning_model_mask = (
        frame["BAG"].eq("Structural")
        & frame["Discovery arm"].eq("Redundancy arm")
    ) | (
        frame["BAG"].eq("Functional")
        & frame["Discovery arm"].eq("Synergy arm")
    )
    diagnosis_families = (
        frame.index[diagnosis_mask & winning_model_mask],
        frame.index[diagnosis_mask & ~winning_model_mask],
    )
    for indices in diagnosis_families:
        frame.loc[indices, "Holm-adjusted p"] = multipletests(
            frame.loc[indices, "P value"].astype(float), method="holm"
        )[1]
    random_indices = frame.index[frame["Analysis"].eq("Random intercept")]
    frame.loc[random_indices, "Holm-adjusted p"] = multipletests(
        frame.loc[random_indices, "P value"].astype(float), method="holm"
    )[1]
    return PublicationTable(12, "Linear mixed-effects models of signed BAG bias", frame, "Models were fitted by maximum likelihood with country as a random intercept. Fixed-effect P values are from two-sided Wald tests. Diagnosis contrasts were Holm-adjusted separately across the four primary contrasts in the BAG-specific winning models (structural redundancy arm and functional synergy arm) and the four secondary contrasts in the alternative-arm models. Intercept and covariate P values are descriptive and were not multiplicity-adjusted. Random-intercept P values are from 10,000-draw parametric-bootstrap likelihood-ratio tests and were Holm-adjusted across the four models.")


def _table_13(source) -> PublicationTable:
    domain = _get(source, "ST13_AltRepresentations", "Domain-balanced representations")
    domain = domain.rename(columns={"Alternative representation": "Representation", "Selected components": "Components", "Alternative R²": "LOCO R²", "ΔR², alternative − main": "ΔR² versus main"})
    domain["Variance explained"] = None
    domain["Matched R², main vs alternative"] = domain.apply(
        lambda row: f"{row['Matched main R²']:.3f} vs {row['Matched alternative R²']:.3f}",
        axis=1,
    )
    domain["Matched ΔR² (95% CI)"] = domain.apply(
        lambda row: (
            f"{row['Matched ΔR², main − alternative']:.3f} "
            f"({row['95% CI, lower']:.3f} to {row['95% CI, upper']:.3f})"
        ),
        axis=1,
    )
    pca = _get(source, "ST13_AltRepresentations", "Whole-exposome PCA summary")
    pca["Representation"] = "Whole-exposome PCA"
    pca = pca.rename(columns={"PCs": "Components", "PCA R²": "LOCO R²", "ΔR², PCA − main": "ΔR² versus main"})
    pca["Variance explained"] = pca.apply(
        lambda row: f"PC1 {row['PC1 variance fraction']:.3f}; cumulative {row['Cumulative variance fraction']:.3f}",
        axis=1,
    )
    inference = pd.read_csv(ROOT / "outputs" / "dedup" / "model_comparison" / "comparison_results.csv")
    inference = inference[
        inference["comparison"].str.endswith("syn_vs_whole_pca|xgb_tree_d3")
        & inference["loss"].eq("sq")
    ].copy()
    inference["BAG"] = inference["comparison"].str.split("|").str[0].str.title()
    inference = inference.set_index("BAG")
    pca["Matched R², main vs alternative"] = pca["BAG"].map(
        lambda bag: f"{inference.loc[bag, 'r2_a']:.3f} vs {inference.loc[bag, 'r2_b']:.3f}"
    )
    pca["Matched ΔR² (95% CI)"] = pca["BAG"].map(
        lambda bag: (
            f"{inference.loc[bag, 'delta_r2']:.3f} "
            f"({inference.loc[bag, 'delta_r2_ci_lo']:.3f} to {inference.loc[bag, 'delta_r2_ci_hi']:.3f})"
        )
    )
    pca["Country-cluster bootstrap p"] = pca["BAG"].map(
        lambda bag: inference.loc[bag, "boot_p_delta_r2"]
    )
    columns = ["BAG", "Representation", "Components", "LOCO R²", "Main-analysis R²", "ΔR² versus main", "Variance explained", "Matched R², main vs alternative", "Matched ΔR² (95% CI)", "Country-cluster bootstrap p"]
    frame = _concat_records([domain, pca], columns)
    frame = _holm_adjust(
        frame,
        "Country-cluster bootstrap p",
        ["BAG"],
        adjusted_column="Holm-adjusted bootstrap p",
    )
    return PublicationTable(13, "Alternative exposome representations", frame, "Principal-component transformations were fitted within each training fold. LOCO R² columns summarize the sensitivity analysis; matched columns compare the deployed depth-3 synergy-arm model with each selected alternative using paired out-of-fold predictions and two-sided country-cluster bootstraps on ΔR² (10,000 draws). P values were Holm-adjusted across the three representations within each BAG measure. Incremental results are provided in Supplementary Fig. S4 Source Data.")


def _table_14(source) -> PublicationTable:
    performance = _get(source, "ST14_NegativeOmega", "Selected-model performance").set_index("BAG").add_prefix("Performance: ")
    pools = _get(source, "ST14_NegativeOmega", "Candidate pools and diversity slopes").set_index("BAG").add_prefix("Candidate analysis: ")
    comparison = _get(source, "ST14_NegativeOmega", "Matched-country arm comparison").set_index("BAG").add_prefix("Matched-country comparison: ")
    frame = pd.concat([performance, pools, comparison], axis=1).T.reset_index().rename(columns={"index": "Metric"})
    frame = frame[["Metric", "Structural", "Functional"]]
    p_row = frame["Metric"].eq("Matched-country comparison: Two-sided p")
    adjusted_values = multipletests(
        frame.loc[p_row, ["Structural", "Functional"]].iloc[0].astype(float),
        method="holm",
    )[1]
    adjusted_row = pd.DataFrame(
        [{
            "Metric": "Matched-country comparison: Holm-adjusted p",
            "Structural": adjusted_values[0],
            "Functional": adjusted_values[1],
        }]
    )
    insertion = frame.index[p_row][0] + 1
    frame = pd.concat(
        [frame.iloc[:insertion], adjusted_row, frame.iloc[insertion:]],
        ignore_index=True,
    )
    return PublicationTable(14, "Negative-O-information arm-definition sensitivity", frame, "The restricted analysis additionally required evaluated O-information below zero for synergy-arm candidates. Matched-country arm comparisons used two-sided paired Wilcoxon signed-rank tests; P values were Holm-adjusted across the two BAG measures. Country-level rows remain in the audit workbook.")


def _table_15(source) -> PublicationTable:
    selected = _get(source, "ST15_DiagnosticContext", "Selected models by diagnostic context")
    selected_inference = pd.read_csv(
        ROOT / "outputs/sensitivity/dedup/normative_context/manuscript_selected_model_inference.csv"
    )
    selected_inference["bag"] = selected_inference["bag"].str.title()
    selected_inference["condition"] = selected_inference["condition"].str.replace("CN", "HC", regex=False).str.replace("FTD", "FTLD", regex=False)
    selected_inference = selected_inference.set_index(["bag", "condition"])
    selected_rows = pd.DataFrame({
        "Analysis": "Selected model",
        "BAG": selected["BAG"],
        "Setting": selected["Diagnostic context"],
        "Model or arm": selected["Selected level"].astype(str) + "; " + selected["Selected arm"].astype(str),
        "N": [f"{int(selected_inference.loc[(bag, setting), 'n_train']):,}/{int(selected_inference.loc[(bag, setting), 'n_test']):,}" for bag, setting in zip(selected["BAG"], selected["Diagnostic context"])],
        "Primary estimate": selected["Selected-model R²"],
        "Comparator estimate": selected["Baseline R²"],
        "Difference or interaction": selected["ΔR²"],
        "P value": [selected_inference.loc[(bag, setting), "country_wilcoxon_p"] for bag, setting in zip(selected["BAG"], selected["Diagnostic context"])],
    })
    diversity = _get(source, "ST15_DiagnosticContext", "Domain-diversity slopes by diagnostic context")
    diversity_rows = pd.DataFrame({
        "Analysis": "Domain diversity",
        "BAG": diversity["BAG"],
        "Setting": diversity["Diagnostic context"],
        "Model or arm": "Redundancy versus synergy arm",
        "N": diversity["Candidates"],
        "Primary estimate": diversity["Redundancy-arm diversity slope"],
        "Comparator estimate": diversity["Synergy-arm diversity slope"],
        "Difference or interaction": diversity["Synergy-arm diversity slope"] - diversity["Redundancy-arm diversity slope"],
        "P value": diversity["Arm-interaction p"],
    })
    weighting = _get(source, "ST15_DiagnosticContext", "Equal-diagnosis-weighting sensitivity")
    weighting_inference = pd.read_csv(
        ROOT / "outputs/sensitivity/dedup/diagnosis_balance/manuscript_weighting_inference.csv"
    )
    weighting_inference["bag"] = weighting_inference["bag"].str.title()
    weighting_inference["objective"] = weighting_inference["objective"].replace({"o_min": "Synergy arm", "o_max": "Redundancy arm"})
    weighting_inference["diagnosis"] = weighting_inference["diagnosis"].replace({"CN": "HC", "FTD": "FTLD"})
    weighting_inference = weighting_inference.set_index(["bag", "objective", "diagnosis"])
    weighting_rows = pd.DataFrame({
        "Analysis": "Equal diagnosis weighting",
        "BAG": weighting["BAG"],
        "Setting": weighting["Diagnosis"],
        "Model or arm": weighting["Discovery arm"],
        "N": weighting["Participants"],
        "Primary estimate": weighting["Weighted mean bias"],
        "Comparator estimate": weighting["Unweighted mean bias"],
        "Difference or interaction": weighting["Δ mean bias"],
        "P value": [weighting_inference.loc[(bag, arm, diagnosis), "country_cluster_bootstrap_p"] for bag, arm, diagnosis in zip(weighting["BAG"], weighting["Discovery arm"], weighting["Diagnosis"])],
    })
    columns = list(selected_rows.columns)
    frame = _concat_records([selected_rows, diversity_rows, weighting_rows], columns)
    frame = _holm_adjust(
        frame,
        "P value",
        ["Analysis", "BAG"],
        adjusted_column="Holm-adjusted p",
    )
    return PublicationTable(15, "Diagnostic-context and diagnosis-weighting sensitivity", frame, "For selected-model rows, N is the total eligible training/test cohort (train/test); the actual training N varies by LOCO fold because the held-out country is excluded. Estimates are selected-model and baseline R²; P values are from two-sided paired Wilcoxon tests across countries. For domain-diversity rows, estimates are redundancy- and synergy-arm slopes and the difference is synergy minus redundancy. For weighting rows, estimates are weighted and unweighted mean bias in years; P values are from two-sided country-cluster bootstraps with 10,000 draws. P values were Holm-adjusted across settings within each analysis and BAG measure.")


def _table_16(source) -> PublicationTable:
    folds = _get(source, "ST16_Geography", "Country and region folds")
    fold_rows = pd.DataFrame({
        "Analysis": "Held-out performance",
        "BAG": folds["BAG"],
        "Unit or term": folds["Held-out unit"],
        "N": folds["Folds"],
        "Estimate": folds["Positive fraction"],
        "SE": None,
        "95% CI or range": folds.apply(lambda row: f"{row['Minimum fold R²']:.3f} to {row['Maximum fold R²']:.3f}", axis=1),
        "P value": None,
        "Additional statistic": folds["Non-positive folds"].fillna("None").map(
            lambda value: f"Non-positive folds: {value}"
        ),
    })
    meta = _ci(_get(source, "ST16_Geography", "Country-performance meta-regression"))
    meta_rows = pd.DataFrame({
        "Analysis": "Country meta-regression",
        "BAG": meta["BAG"],
        "Unit or term": meta["Term"],
        "N": meta["Countries"],
        "Estimate": meta["Estimate"],
        "SE": meta["SE"],
        "95% CI or range": meta["95% CI"],
        "P value": meta["Two-sided p"],
        "Additional statistic": meta["Model R²"].map(lambda value: f"Model R²={value:.3f}"),
    })
    frame = _concat_records([fold_rows, meta_rows], list(fold_rows.columns))
    frame["Holm-adjusted p"] = None
    meta_mask = frame["Analysis"].eq("Country meta-regression")
    for indices in frame.loc[meta_mask].groupby("BAG", sort=False).groups.values():
        frame.loc[indices, "Holm-adjusted p"] = multipletests(
            frame.loc[indices, "P value"].astype(float), method="holm"
        )[1]
    return PublicationTable(16, "Geographical generalisation", frame, "Held-out performance uses the depth-3 synergy-arm series and reports the fraction of country or region folds with positive R² and the full fold range. Country meta-regressions model country-level out-of-fold R²; coefficient P values were Holm-adjusted across the three displayed terms within each BAG measure.")


def _table_17(source) -> PublicationTable:
    inference = pd.read_csv(
        ROOT / "outputs/sensitivity/dedup/covariate_target_matched_inference.csv"
    )
    inference["bag"] = inference["bag"].str.title()
    inference = inference.set_index(
        ["analysis", "bag", "specification_or_model"]
    )

    def p_value(analysis: str, bag: str, specification: str):
        value = inference.loc[
            (analysis, bag, specification), "paired_country_wilcoxon_p"
        ]
        return None if pd.isna(value) else value

    covariates = _get(source, "ST17_CovariatesTargets", "Education and scanner covariates")
    covariate_rows = pd.DataFrame({
        "Analysis": "Additional covariates",
        "BAG": covariates["BAG"],
        "Specification or model": covariates["Covariate specification"],
        "Participants": covariates["Participants"],
        "LOCO R²": covariates["LOCO R²"],
        "Reference R²": covariates["Reference R²"],
        "ΔR²": covariates["ΔR² versus reference"],
        "P value": [p_value("Additional covariates", bag, specification) for bag, specification in zip(covariates["BAG"], covariates["Covariate specification"])],
    })
    residual = _get(source, "ST17_CovariatesTargets", "Residualized BAG target")
    rows = []
    mapping = {
        "Covariate baseline": ("Baseline R²", None),
        "Best single exposure": ("Best single R²", "Single ΔR²"),
        "Best synergy-arm model": ("Best synergy-arm R²", "Synergy ΔR²"),
        "Best redundancy-arm model": ("Best redundancy-arm R²", "Redundancy ΔR²"),
    }
    for _, row in residual.iterrows():
        for model, (r2_column, delta_column) in mapping.items():
            rows.append({
                "Analysis": "Residualised BAG target",
                "BAG": row["BAG"],
                "Specification or model": model,
                "Participants": None,
                "LOCO R²": row[r2_column],
                "Reference R²": row["Baseline R²"],
                "ΔR²": 0.0 if delta_column is None else row[delta_column],
                "P value": p_value("Residualised BAG target", row["BAG"], model),
            })
    frame = _concat_records([covariate_rows, pd.DataFrame(rows)], list(covariate_rows.columns))
    frame = _holm_adjust(
        frame,
        "P value",
        ["Analysis", "BAG"],
        adjusted_column="Holm-adjusted p",
    )
    return PublicationTable(17, "Covariate and BAG-target sensitivity analyses", frame, "Additional-covariate comparisons used the same complete-case sample. In the residualised-target analysis, age, sex and diagnosis effects were fitted within each training fold and applied to held-out participants. P values compare each specification with its displayed reference using two-sided paired Wilcoxon tests across countries and were Holm-adjusted across non-reference specifications within each analysis and BAG measure; reference rows are not tested against themselves.")


def _table_18(source) -> PublicationTable:
    selected = _get(source, "ST18_SetSizeCaps", "Selected model by maximum set size")
    diversity = _get(source, "ST18_SetSizeCaps", "Domain-diversity effect by maximum set size")
    caps = {5, 10, 15, 20, 21, 22, 25, 30}
    selected = selected[selected["Maximum set size"].isin(caps)][["BAG", "Maximum set size", "Best R²", "Selected set size", "Selected arm"]]
    diversity = diversity[diversity["Maximum set size"].isin(caps)][["BAG", "Maximum set size", "Redundancy-arm slope", "Synergy-arm slope", "Arm interaction", "Within-order permutation p"]]
    frame = selected.merge(diversity, on=["BAG", "Maximum set size"], validate="one_to_one")
    frame = _holm_adjust(
        frame,
        "Within-order permutation p",
        ["BAG"],
        adjusted_column="Holm-adjusted permutation p",
    )
    return PublicationTable(18, "Maximum candidate-set-size sensitivity", frame, "Representative caps include the prespecified cap 21, the sign-change boundary at cap 22 and the final cap 30. P values were Holm-adjusted across the eight displayed caps within each BAG measure. Results for every cap from 5 to 30 are provided in Supplementary Fig. S10 Source Data.")


BUILDERS: tuple[Callable, ...] = (
    _table_01, _table_02, _table_03, _table_04, _table_05, _table_06,
    _table_07, _table_08, _table_09, _table_10, _table_11, _table_12,
    _table_13, _table_14, _table_15, _table_16, _table_17, _table_18,
)


def _number_format(header: str) -> str:
    lower = header.lower()
    if "p value" in lower or lower.endswith(" p"):
        return "0.0000E+00"
    if any(token in lower for token in ("participants", "countries", "candidates", "models", "set size", "folds", "combinations")) and not any(token in lower for token in ("fraction", "mean", "median", "r²")):
        return "0"
    if any(token in lower for token in ("r²", "estimate", "slope", "difference", "bias", "mae", "rmse", "fraction", "ratio", "icc", "variance", "Ω", "entropy", "cohen", "effect", "se", "statistic", "pearson")):
        return "0.000"
    return "General"


def _write_table(workbook: Workbook, table: PublicationTable) -> None:
    worksheet = workbook.create_sheet(table.sheet_name)
    worksheet.sheet_view.showGridLines = False
    columns = len(table.frame.columns)
    worksheet.cell(1, 1, f"{table.label} | {table.title}")
    worksheet.cell(1, 1).font = Font(name=FONT, size=11, bold=True, color=BLACK)
    worksheet.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    worksheet.row_dimensions[1].height = 30

    for index, header in enumerate(table.frame.columns, start=1):
        cell = worksheet.cell(3, index, header)
        cell.font = Font(name=FONT, size=10, bold=True, color=BLACK)
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        cell.border = Border(top=RULE, bottom=RULE)
    for row_index, values in enumerate(table.frame.itertuples(index=False, name=None), start=4):
        for column_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row_index, column_index, value)
            cell.font = Font(name=FONT, size=10, color=BLACK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, float):
                cell.number_format = _number_format(str(table.frame.columns[column_index - 1]))

    final_data_row = 3 + len(table.frame)
    for column in range(1, columns + 1):
        worksheet.cell(final_data_row, column).border = Border(bottom=RULE)
    note_row = final_data_row + 2
    worksheet.cell(note_row, 1, "Note. " + table.note)
    worksheet.cell(note_row, 1).font = Font(name=FONT, size=9, color=GREY)
    worksheet.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=columns)
    worksheet.row_dimensions[note_row].height = max(30, min(75, 15 * (1 + len(table.note) // 140)))

    for column_index, header in enumerate(table.frame.columns, start=1):
        values = [str(header)] + [str(value or "") for value in table.frame.iloc[:, column_index - 1]]
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(24, max(9, max(map(len, values)) + 1))
    worksheet.freeze_panes = "A4"
    worksheet.print_title_rows = "3:3"
    worksheet.print_area = f"A1:{get_column_letter(columns)}{note_row}"
    worksheet.page_setup.orientation = "portrait"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = worksheet.page_margins.right = 0.35
    worksheet.page_margins.top = worksheet.page_margins.bottom = 0.5


def build(output: Path = OUTPUT) -> Path:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Canonical supplementary workbook is missing: {SOURCE}")
    source = load_workbook(SOURCE, data_only=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    tables = [builder(source) for builder in BUILDERS]
    assert [table.number for table in tables] == list(range(1, 19))
    for table in tables:
        _write_table(workbook, table)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def main() -> None:
    print(build())


if __name__ == "__main__":
    main()
