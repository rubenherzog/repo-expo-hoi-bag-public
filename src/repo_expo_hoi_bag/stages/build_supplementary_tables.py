"""Build the manuscript-ready supplementary-table workbook.

The workbook contains the 18 supplementary tables cited by
``docs/MS_Expo_HOI_BAG_results_methods.md``.  It only aggregates existing
analysis outputs; it does not fit or refit any statistical model.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
OUTPUT = ROOT / "docs" / "supplementary" / "Supplementary_Tables.xlsx"

NAVY = "1F4E78"
BLUE = "D9EAF7"
PALE = "EAF2F8"
WHITE = "FFFFFF"
GREY = "666666"
THIN_GREY = Side(style="thin", color="B7B7B7")


def read_csv(path: str) -> pd.DataFrame:
    """Read a repository-relative CSV and fail with a useful path."""
    full_path = ROOT / path
    if not full_path.exists():
        raise FileNotFoundError(f"Required supplementary-table source is missing: {full_path}")
    return pd.read_csv(full_path)


def publication_frame(frame: pd.DataFrame, columns: dict[str, str]) -> pd.DataFrame:
    """Select and rename columns for a publication-facing worksheet."""
    missing = set(columns).difference(frame.columns)
    if missing:
        raise ValueError(f"Missing columns {sorted(missing)} from {list(frame.columns)}")
    result = frame.loc[:, list(columns)].rename(columns=columns).reset_index(drop=True)
    if "BAG" in result.columns:
        bag_order = {"structural": 0, "functional": 1}
        order = result["BAG"].astype(str).str.lower().map(bag_order).fillna(2)
        result = result.assign(_bag_order=order).sort_values("_bag_order", kind="stable")
        result = result.drop(columns="_bag_order").reset_index(drop=True)
    return result


def arm_labels(series: pd.Series) -> pd.Series:
    return series.replace(
        {
            "o_min": "Synergy arm",
            "o_max": "Redundancy arm",
            "syn": "Synergy arm",
            "red": "Redundancy arm",
        }
    )


def diagnosis_labels(series: pd.Series) -> pd.Series:
    return series.replace({"CN": "HC", "FTD": "FTLD", "AD-CN": "AD-HC", "FTD-CN": "FTLD-HC"})


class TableSheet:
    """Small helper for consistently formatted, copyable Nature tables."""

    def __init__(self, workbook: Workbook, name: str, number: int, title: str, note: str):
        self.ws = workbook.create_sheet(name)
        self.row = 1
        self.ws.sheet_view.showGridLines = False
        self.ws.cell(self.row, 1, f"Supplementary Table {number} | {title}")
        self.ws.cell(self.row, 1).font = Font(bold=True, size=14, color=WHITE)
        self.ws.cell(self.row, 1).fill = PatternFill("solid", fgColor=NAVY)
        self.ws.cell(self.row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=12)
        self.ws.row_dimensions[self.row].height = 30
        self.row += 2
        self.note(note)

    def note(self, text: str) -> None:
        self.ws.cell(self.row, 1, text)
        self.ws.cell(self.row, 1).font = Font(italic=True, color=GREY, size=10)
        self.ws.cell(self.row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=12)
        self.ws.row_dimensions[self.row].height = max(30, min(90, 15 * (1 + len(text) // 130)))
        self.row += 2

    def section(self, heading: str, frame: pd.DataFrame) -> None:
        if "BAG" in frame.columns:
            bag_order = {"structural": 0, "functional": 1}
            order = frame["BAG"].astype(str).str.lower().map(bag_order).fillna(2)
            frame = (
                frame.assign(_bag_order=order)
                .sort_values("_bag_order", kind="stable")
                .drop(columns="_bag_order")
                .reset_index(drop=True)
            )
        self.ws.cell(self.row, 1, heading)
        self.ws.cell(self.row, 1).font = Font(bold=True, size=11)
        self.ws.cell(self.row, 1).fill = PatternFill("solid", fgColor=BLUE)
        self.ws.merge_cells(
            start_row=self.row,
            start_column=1,
            end_row=self.row,
            end_column=max(1, len(frame.columns)),
        )
        self.row += 1
        header_row = self.row
        for column_index, header in enumerate(frame.columns, start=1):
            cell = self.ws.cell(self.row, column_index, str(header))
            cell.font = Font(bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
        self.row += 1
        for values in frame.itertuples(index=False, name=None):
            for column_index, value in enumerate(values, start=1):
                if pd.isna(value):
                    value = None
                elif isinstance(value, np.generic):
                    value = value.item()
                cell = self.ws.cell(self.row, column_index, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = Border(bottom=THIN_GREY)
                if isinstance(value, float):
                    cell.number_format = "0.000000"
            self.row += 1
        self.ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(max(1, len(frame.columns)))}{self.row - 1}"
        )
        self.row += 2

    def finish(self) -> None:
        if self.ws.max_column > 12:
            self.ws.unmerge_cells("A1:L1")
            self.ws.merge_cells(
                start_row=1, start_column=1, end_row=1, end_column=self.ws.max_column
            )
        self.ws.freeze_panes = "A5"
        self.ws.page_setup.orientation = "landscape"
        self.ws.page_setup.fitToWidth = 1
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        for column in range(1, self.ws.max_column + 1):
            values: Iterable[str] = (
                str(self.ws.cell(row, column).value or "")
                for row in range(1, self.ws.max_row + 1)
            )
            width = min(45, max(11, max((len(value) for value in values), default=0) + 2))
            self.ws.column_dimensions[get_column_letter(column)].width = width


def table_01(workbook: Workbook) -> None:
    base = "outputs/figures/dedup/paper/complete/source_data/fig2_grid_v3_max_30_dedup_source_data_"
    omega = read_csv(base + "a1_oinfo_by_set_size.csv")
    domains = read_csv(base + "a2_unique_domains.csv")
    entropy = read_csv(base + "a3_shannon_entropy.csv")
    order = omega.merge(domains, on=["objective", "set_size"]).merge(
        entropy, on=["objective", "set_size"]
    )
    greedy = read_csv("outputs/dedup/greedy_synergy_by_order.csv").rename(
        columns={"order": "set_size"}
    )
    order = order.merge(greedy[["set_size", "syn_median"]], on="set_size", how="left")
    order["omega_median"] = np.where(order["objective"] == "o_min", order["syn_median"], np.nan)
    order["arm"] = arm_labels(order["objective"])
    order["domain_sd"] = order["n_domains_plus_1sd"] - order["n_domains_mean"]
    order["entropy_sd"] = order["shannon_h_plus_1sd"] - order["shannon_h_mean"]
    order = publication_frame(
        order.sort_values(["set_size", "objective"]),
        {
            "set_size": "Set size",
            "arm": "Discovery arm",
            "omega_min": "Minimum Ω (nats)",
            "omega_median": "Median Ω (nats)",
            "omega_max": "Maximum Ω (nats)",
            "n_domains_mean": "Domains, mean",
            "domain_sd": "Domains, SD",
            "shannon_h_mean": "Shannon H, mean (bits)",
            "entropy_sd": "Shannon H, SD (bits)",
        },
    )
    baseline = publication_frame(
        read_csv("outputs/dedup/subcomb_oinfo/subcomb_baseline_stats.csv"),
        {
            "order_k": "Combination order",
            "n_total": "Total combinations",
            "n_neg": "Negative-Ω combinations",
            "frac_neg": "Negative-Ω fraction",
            "ci_lo": "95% CI, lower",
            "ci_hi": "95% CI, upper",
        },
    )
    regression = publication_frame(
        read_csv("outputs/dedup/subcomb_oinfo/subcomb_ols_coefficients.csv"),
        {
            "bag": "BAG",
            "order_k": "Combination order",
            "param": "Term",
            "beta": "Estimate",
            "se": "SE",
            "ci_lo": "95% CI, lower",
            "ci_hi": "95% CI, upper",
            "p_permutation": "Permutation p",
            "perm_n": "Permutations",
            "n": "Candidates",
        },
    )
    sheet = TableSheet(
        workbook,
        "ST01_HigherOrder",
        1,
        "Higher-order information structure and synergy baselines",
        "Candidate sets were discovered on 259 unique country-year signatures. Arms denote the optimisation objective. Values at each set size summarize the retained candidate pool; H is Shannon domain entropy.",
    )
    sheet.section("Information structure by set size", order)
    sheet.section("Exposome-wide negative-Ω baselines", baseline)
    sheet.section("Sub-combination enrichment regressions", regression)
    sheet.note("Permutation tests are two-sided within model level with 10,000 label permutations. Classical OLS p values are omitted because overlapping candidate sets are not independent observations.")
    sheet.finish()


def table_02(workbook: Workbook) -> None:
    keep = read_csv(
        "outputs/dedup/model_comparison/revision_inference/model_complexity_best_per_level.csv"
    )
    keep["bag"] = keep["bag"].str.title()
    keep["objective"] = arm_labels(keep["objective"])
    keep["rung_id"] = keep["rung_id"].replace(
        {"ols": "OLS", "xgb_tree_d1": "d1", "xgb_tree_d2": "d2", "xgb_tree_d3": "d3"}
    )
    table = publication_frame(
        keep,
        {
            "bag": "BAG",
            "objective": "Discovery arm",
            "rung_id": "Model level",
            "full_r2": "Best LOCO R²",
            "order": "Selected set size",
        },
    )
    single = read_csv(
        "outputs/dedup/model_comparison/revision_inference/single_exposure_complexity_per_level.csv"
    )
    single["bag"] = single["bag"].str.title()
    single["rung_id"] = single["rung_id"].replace(
        {"ols": "OLS", "xgb_tree_d1": "d1", "xgb_tree_d2": "d2", "xgb_tree_d3": "d3"}
    )
    single = publication_frame(
        single,
        {
            "bag": "BAG",
            "rung_id": "Model level",
            "selected_exposure": "Selected exposure",
            "full_r2": "Best single-exposure LOCO R²",
            "n_subjects": "Participants",
            "n_countries": "Countries",
        },
    )
    paired = read_csv(
        "outputs/dedup/model_comparison/revision_inference/single_exposure_complexity_paired.csv"
    )
    paired["bag"] = paired["bag"].str.title()
    paired = publication_frame(
        paired,
        {
            "bag": "BAG",
            "n_subjects": "Participants",
            "n_countries": "Countries",
            "r2_d3": "Best d3 R²",
            "r2_ols": "Best OLS R²",
            "delta_r2": "ΔR², d3 − OLS",
            "delta_r2_ci_lo": "95% CI, lower",
            "delta_r2_ci_hi": "95% CI, upper",
            "country_cluster_bootstrap_p": "Country-cluster bootstrap p",
            "bootstrap_draws": "Bootstrap draws",
        },
    )
    sheet = TableSheet(
        workbook,
        "ST02_ModelComplexity",
        2,
        "Effects of model complexity",
        "The best candidate or single exposure was selected independently within each BAG and model level. Values are pooled out-of-fold R². The paired contrast compares the selected d3 and OLS single-exposure predictions by resampling countries.",
    )
    sheet.section("Best multivariate candidate at each model level", table)
    sheet.section("Best single exposure at each model level", single)
    sheet.section("Paired d3-versus-OLS single-exposure comparison", paired)
    sheet.note("Multivariate winner trajectories are descriptive because a different candidate could be selected at each level. The single-exposure contrast uses matched out-of-fold predictions and a two-sided country-cluster bootstrap on ΔR² (10,000 draws); because the selected exposure can differ by level, it does not isolate a within-exposure complexity effect.")
    sheet.finish()


def table_03(workbook: Workbook) -> None:
    comparisons = read_csv("outputs/dedup/model_comparison/comparison_results.csv")
    selected_names = [
        "best_syn_vs_baseline",
        "best_red_vs_baseline",
        "best_single_vs_baseline",
        "syn_vs_red",
    ]
    pattern = "|".join(selected_names)
    selected = comparisons[
        (comparisons["loss"] == "sq")
        & comparisons["comparison"].str.contains(pattern)
        & comparisons["comparison"].str.endswith("xgb_tree_d3")
    ].copy()
    selected["BAG"] = selected["comparison"].str.split("|").str[0].str.title()
    selected["Comparison"] = (
        selected["comparison"]
        .str.split("|").str[1]
        .replace(
            {
                "best_syn_vs_baseline": "Best synergy arm vs covariate baseline",
                "best_red_vs_baseline": "Best redundancy arm vs covariate baseline",
                "best_single_vs_baseline": "Best single exposure vs covariate baseline",
                "syn_vs_red": "Best synergy arm vs best redundancy arm",
            }
        )
    )
    tests = publication_frame(
        selected.sort_values(["BAG", "Comparison"]),
        {
            "BAG": "BAG",
            "Comparison": "Comparison (A vs B)",
            "n_subjects": "Participants",
            "n_countries": "Countries",
            "r2_a": "R², model A",
            "r2_b": "R², model B",
            "delta_r2": "ΔR², A − B",
            "delta_r2_ci_lo": "95% CI, lower",
            "delta_r2_ci_hi": "95% CI, upper",
            "boot_p_delta_r2": "Country-cluster bootstrap p",
        },
    )
    multivariate_vs_single = read_csv(
        "outputs/dedup/model_comparison/revision_inference/best_multivariate_vs_single.csv"
    )
    multivariate_vs_single["bag"] = multivariate_vs_single["bag"].str.title()
    multivariate_vs_single["comparison"] = (
        "Best multivariate model ("
        + multivariate_vs_single["selected_arm"].str.lower()
        + ") vs best single exposure"
    )
    multivariate_vs_single = publication_frame(
        multivariate_vs_single,
        {
            "bag": "BAG",
            "comparison": "Comparison (A vs B)",
            "n_subjects": "Participants",
            "n_countries": "Countries",
            "r2_multivariate": "R², model A",
            "r2_single": "R², model B",
            "delta_r2": "ΔR², A − B",
            "delta_r2_ci_lo": "95% CI, lower",
            "delta_r2_ci_hi": "95% CI, upper",
            "country_cluster_bootstrap_p": "Country-cluster bootstrap p",
        },
    )
    tests = pd.concat([tests, multivariate_vs_single], ignore_index=True)
    consistency = read_csv("outputs/dedup/model_comparison/single_vs_best_country_consistency.csv")
    consistency = consistency[consistency["loss"] == "abs"].copy()
    consistency["bag"] = consistency["bag"].str.title()
    consistency["objective"] = arm_labels(consistency["objective"])
    consistency = publication_frame(
        consistency,
        {
            "bag": "BAG",
            "objective": "High-order arm",
            "n_subjects": "Participants",
            "n_countries": "Countries",
            "country_mean_d": "Country-mean loss difference",
            "country_t_p": "Country-mean t-test p",
            "ols_beta0": "Cluster-robust estimate",
            "ols_cluster_se": "Cluster-robust SE",
            "ols_cluster_p": "Cluster-robust p",
            "win_fraction": "Country win fraction",
            "win_frac_ci_lo": "Win fraction 95% CI, lower",
            "win_frac_ci_hi": "Win fraction 95% CI, upper",
            "weighted_t_p": "Weighted-country t-test p",
        },
    )
    effects = tests[tests["Comparison (A vs B)"].str.contains("baseline")].copy()
    effects["Cohen f²"] = (effects["R², model A"] - effects["R², model B"]) / (
        1 - effects["R², model A"]
    )
    effects = effects[["BAG", "Comparison (A vs B)", "Cohen f²"]]
    sheet = TableSheet(
        workbook,
        "ST03_BestModels",
        3,
        "Best multivariate and single-exposure models",
        "All headline comparisons use matched depth-3 out-of-fold predictions and the corresponding depth-3 covariate baseline. This avoids mixing baseline values from different model levels or analysis bundles.",
    )
    sheet.section("Matched depth-3 R² comparisons", tests)
    sheet.section("Incremental effect sizes versus covariate baseline", effects)
    sheet.section("Country-consistency analyses: arm-specific multivariate models versus best single exposure", consistency)
    sheet.note("Difference tests use two-sided country-cluster bootstraps with 10,000 iterations. The headline multivariate-versus-single comparison uses the overall depth-3 arm winner for each BAG: the redundancy-arm winner for structural BAG and the synergy-arm winner for functional BAG.")
    sheet.finish()


def table_04(workbook: Workbook) -> None:
    frames = []
    for bag in ("structural", "functional"):
        frame = read_csv(f"outputs/sensitivity/dedup/country_block_null/{bag}_winner_pvalues.csv")
        frames.append(frame[frame["rung_id"] == "xgb_tree_d3"])
    data = pd.concat(frames, ignore_index=True)
    data["bag"] = data["bag"].str.title()
    data["label"] = data["label"].replace(
        {"best_syn": "Best synergy arm", "best_red": "Best redundancy arm", "best_single": "Best single exposure"}
    )
    table = publication_frame(
        data,
        {
            "bag": "BAG",
            "label": "Model",
            "order": "Set size",
            "observed_r2": "Observed R²",
            "null_mean": "Null mean R²",
            "null_std": "Null SD",
            "cohen_d_vs_null": "z versus null",
            "p_value": "Empirical upper-tail p",
            "n_perm": "Permutations",
        },
    )
    sheet = TableSheet(workbook, "ST04_CountryBlockNull", 4, "Country-block permutation null and selection sensitivity", "Country-year exposure signatures were permuted as blocks while participants, BAG, covariates and LOCO folds were held fixed. Empirical p=(k+1)/(n+1).")
    sheet.section("Depth-3 deployed models", table)
    crossfit = read_csv(
        "outputs/dedup/model_comparison/revision_inference/country_crossfit_selection_summary.csv"
    )
    crossfit["bag"] = crossfit["bag"].str.title()
    crossfit = publication_frame(
        crossfit,
        {
            "bag": "BAG",
            "arm": "Discovery arm",
            "held_out_countries": "Held-out countries",
            "unique_selected_candidates": "Distinct selected candidates",
            "most_frequent_candidate_fraction": "Most frequent candidate fraction",
            "median_held_out_country_r2": "Median held-out-country R²",
            "q1_held_out_country_r2": "R², first quartile",
            "q3_held_out_country_r2": "R², third quartile",
            "positive_r2_countries": "Countries with R² > 0",
        },
    )
    sheet.section("Country-cross-fitted candidate-selection sensitivity", crossfit)
    crossfit_folds = read_csv(
        "outputs/dedup/model_comparison/revision_inference/country_crossfit_selection_folds.csv"
    )
    crossfit_folds["bag"] = crossfit_folds["bag"].str.title()
    crossfit_folds = publication_frame(
        crossfit_folds,
        {
            "bag": "BAG",
            "arm": "Discovery arm",
            "held_out_country": "Held-out country",
            "selected_set_size": "Selected set size",
            "training_country_rmse": "Other-country RMSE",
            "held_out_country_r2": "Held-out-country R²",
            "held_out_country_rmse": "Held-out-country RMSE",
            "held_out_n": "Held-out participants",
        },
    )
    sheet.section("Country-cross-fitted selections by held-out country", crossfit_folds)
    sheet.note("For each held-out country, the depth-3 candidate was selected by minimum participant-weighted RMSE across all other countries, then evaluated in the unused country. This sensitivity addresses winner-selection stability without claiming to replace the country-block null.")
    sheet.finish()


def top_model_composition() -> tuple[pd.DataFrame, pd.DataFrame]:
    base = ROOT / "outputs/figures/dedup/paper/complete/source_data"
    subcomb = read_csv("outputs/dedup/subcomb_oinfo/subcomb_enrichment.csv")
    summaries = []
    candidate_rows = []
    for bag, prefix, short in (("Structural", "b", "struct"), ("Functional", "c", "func")):
        perf = pd.read_csv(base / f"fig2_grid_v3_max_30_dedup_source_data_{prefix}1_{short}_r2.csv")
        sizes = pd.read_csv(base / f"fig2_grid_v3_max_30_dedup_source_data_{prefix}2_{short}_set_size.csv")
        top = perf.merge(
            sizes[["model_level", "objective", "candidate_id", "order"]],
            on=["model_level", "objective", "candidate_id"],
        )
        top = top[top["model_level"] == "d3"].copy()
        fractions = subcomb[
            (subcomb["bag"] == bag.lower())
            & (subcomb["rung_id"] == "xgb_tree_d3")
            & (subcomb["order_k"] == 3)
        ][["candidate_id", "frac_neg"]]
        top = top.merge(fractions, on="candidate_id", how="left", validate="one_to_one")
        syn = top[top["objective"] == "o_min"]
        red = top[top["objective"] == "o_max"]
        permutation = read_csv(
            "outputs/dedup/model_comparison/revision_inference/overlap_aware_top50_permutation.csv"
        )
        summaries.append(
            {
                "BAG": bag,
                "Metric": "Set size",
                "Synergy-arm median": syn["order"].median(),
                "Redundancy-arm median": red["order"].median(),
                "Two-sided label-permutation p": permutation.loc[
                    (permutation["bag"] == bag.lower())
                    & (permutation["metric"] == "Median set-size difference"),
                    "permutation_p_two_sided",
                ].iloc[0],
                "Models per arm": len(syn),
            }
        )
        summaries.append(
            {
                "BAG": bag,
                "Metric": "Negative-Ω triplet fraction",
                "Synergy-arm median": syn["frac_neg"].median(),
                "Redundancy-arm median": red["frac_neg"].median(),
                "Two-sided label-permutation p": permutation.loc[
                    (permutation["bag"] == bag.lower())
                    & permutation["metric"].str.contains("triplet-fraction"),
                    "permutation_p_two_sided",
                ].iloc[0],
                "Models per arm": len(syn),
            }
        )
        top.insert(0, "BAG", bag)
        top["objective"] = arm_labels(top["objective"])
        candidate_rows.append(top)
    candidates = pd.concat(candidate_rows, ignore_index=True)
    candidates["candidate_rank"] = (
        candidates.groupby(["BAG", "objective"])["full_r2"]
        .rank(method="first", ascending=False)
        .astype(int)
    )
    candidates = publication_frame(
        candidates,
        {
            "BAG": "BAG",
            "objective": "Discovery arm",
            "candidate_rank": "Within-arm performance rank",
            "full_r2": "LOCO R²",
            "order": "Set size",
            "frac_neg": "Negative-Ω triplet fraction",
        },
    )
    return pd.DataFrame(summaries), candidates


def table_05(workbook: Workbook) -> None:
    summary, candidates = top_model_composition()
    permutation = read_csv(
        "outputs/dedup/model_comparison/revision_inference/overlap_aware_top50_permutation.csv"
    )
    permutation["bag"] = permutation["bag"].str.title()
    permutation = publication_frame(
        permutation,
        {
            "bag": "BAG",
            "contrast": "Contrast",
            "metric": "Metric",
            "observed_difference": "Observed difference",
            "null_mean": "Permutation-null mean",
            "null_sd": "Permutation-null SD",
            "permutation_p_two_sided": "Two-sided permutation p",
            "permutation_exceedances": "Exceedances",
            "permutations": "Permutations",
        },
    )
    report = read_csv("outputs/dedup/report_stats/report_stats.csv").set_index("key")
    enrichment_rows = []
    for bag in ("structural", "functional"):
        row = report.loc[f"enrichment.{bag}.top10pct"]
        detail = str(row["detail"])
        ratio = float(detail.split("ratio=")[1].split()[0])
        top_k = int(detail.split("top-k=")[1].split(")")[0])
        fraction = float(row["value"])
        enrichment_rows.append(
            {
                "BAG": bag.title(),
                "Top fraction": 0.10,
                "Models": top_k,
                "Synergy-arm models": round(fraction * top_k),
                "Synergy-arm fraction": fraction,
                "Candidate-pool fraction": 0.5,
                "Enrichment ratio": ratio,
            }
        )
    enrichment = pd.DataFrame(enrichment_rows)
    sheet = TableSheet(workbook, "ST05_TopComposition", 5, "Composition of top-performing models", "The top 50 depth-3 candidates were selected separately within each arm and BAG. Frontier enrichment is descriptive because greedy candidates overlap and are not independent draws.")
    sheet.section("Set size and synergistic-triplet content", summary)
    sheet.section("Overlap-aware arm-label permutation results", permutation)
    sheet.section("Top-frontier arm enrichment", enrichment)
    sheet.section("Top-50 candidate-level data", candidates)
    sheet.note("Set-size and triplet-fraction contrasts use 10,000 two-sided arm-label permutations of the complete selected candidate pool, retaining the candidates, their feature overlap and their joint set-size/triplet structure. No significance test is attached to the descriptive frontier enrichment.")
    sheet.finish()


def table_06(workbook: Workbook) -> None:
    ranges = read_csv(
        "outputs/dedup/model_comparison/revision_inference/diversity_performance_ranges_d3.csv"
    )
    ranges["bag"] = ranges["bag"].str.title()
    ranges["objective"] = arm_labels(ranges["objective"])
    ranges = publication_frame(
        ranges,
        {
            "bag": "BAG",
            "objective": "Discovery arm",
            "n_candidates": "Candidates",
            "minimum_loco_r2": "Minimum LOCO R²",
            "maximum_loco_r2": "Maximum LOCO R²",
        },
    )
    correlations = read_csv("outputs/dedup/model_comparison/diversity_correlations_d3.csv")
    correlations["bag"] = correlations["bag"].str.title()
    correlations["arm"] = correlations["arm"].str.title() + " arm"
    correlations = publication_frame(
        correlations,
        {"bag": "BAG", "arm": "Discovery arm", "n": "Candidates", "pearson_r": "Pearson r (descriptive)", "slope_r2_per_bit": "R² slope per bit (descriptive)"},
    )
    regression = read_csv("outputs/dedup/model_comparison/diversity_regression_d3.csv")
    regression["bag"] = regression["bag"].str.title()
    regression = publication_frame(
        regression,
        {"bag": "BAG", "term": "Estimate", "estimate": "Coefficient", "ci_lo": "95% CI, lower", "ci_hi": "95% CI, upper", "p_cluster_order": "Set-size-clustered p", "n": "Candidates", "n_clusters": "Set-size clusters"},
    )
    permutation = read_csv("outputs/sensitivity/dedup/order_cap/manuscript_diversity_cap_comparison.csv")
    permutation = permutation[permutation["order_cap"] == 30].copy()
    permutation["bag"] = permutation["bag"].str.title()
    permutation = publication_frame(permutation, {"bag": "BAG", "beta_h_red": "Redundancy-arm slope", "beta_h_syn": "Synergy-arm slope", "beta_interaction": "Synergy − redundancy interaction", "p_h_syn_interaction_perm": "Within-order permutation p", "model_r2": "Model R²", "n_obs": "Candidates"})
    sheet = TableSheet(workbook, "ST06_DomainDiversity", 6, "Domain diversity and predictive performance", "Depth-3 candidates, set size 3–30. Pearson coefficients are descriptive. The primary arm-difference inference is the within-order label permutation, which retains the overlapping candidate pool.")
    sheet.section("Observed candidate-level performance ranges", ranges)
    sheet.section("Arm-specific Pearson correlations", correlations)
    sheet.section("Adjusted diversity regression", regression)
    sheet.section("Interaction permutation test", permutation)
    sheet.finish()


def table_07(workbook: Workbook) -> None:
    data = read_csv("outputs/dedup/model_comparison/domain_composition_top20.csv")
    data["bag"] = data["bag"].str.title()
    data["arm"] = data["arm"].str.title() + " arm"
    table = publication_frame(data, {"bag": "BAG", "arm": "Discovery arm", "domain": "Exposome domain", "n_features": "Feature occurrences", "pct_features": "Feature occurrences (%)", "n_models_with_domain": "Models containing domain", "pct_models": "Models containing domain (%)", "top_k": "Models per arm"})
    sheet = TableSheet(workbook, "ST07_DomainComposition", 7, "Domain composition of top-20 models", "Domain composition of the top-20 depth-3 candidates within each BAG and discovery arm. Percentages use all feature occurrences in the corresponding 20 models as denominator.")
    sheet.section("Domain counts and percentages", table)
    sheet.finish()


def table_08(workbook: Workbook) -> None:
    data = read_csv("outputs/dedup/model_comparison/cooccurrence_overlap_permutation.csv")
    data["bag"] = data["bag"].str.title()
    table = publication_frame(data, {"bag": "BAG", "statistic": "Network statistic", "syn": "Synergy arm", "red": "Redundancy arm", "syn_minus_red": "Synergy − redundancy", "null_mean": "Permutation-null mean", "perm_p_two_sided": "Two-sided permutation p", "n_perm": "Permutations"})
    sheet = TableSheet(workbook, "ST08_Cooccurrence", 8, "Domain co-occurrence network architecture", "Top-20 depth-3 candidates per BAG and arm. Arm-label permutations preserve the overlapping candidate pool; empirical p=(k+1)/(n+1).")
    sheet.section("Observed network contrasts", table)
    sheet.finish()


def table_09(workbook: Workbook) -> None:
    top = read_csv("outputs/dedup/model_comparison/recurrent_triplets_top.csv")
    top["bag"] = top["bag"].str.title(); top["arm"] = top["arm"].str.title() + " arm"
    top = publication_frame(top, {"bag": "BAG", "arm": "Discovery arm", "triplet": "Triplet", "n_candidates_with_triplet": "Candidates containing triplet", "n_candidates_in_arm": "Candidates in arm", "prevalence_pct": "Prevalence (%)", "omega": "Ω (nats)", "n_triplets_tied_at_this_prevalence": "Tied triplets"})
    single = read_csv("outputs/dedup/model_comparison/recurrent_triplets_single_exposure.csv")
    single["bag"] = single["bag"].str.title(); single["arm"] = single["arm"].str.title() + " arm"
    single = publication_frame(single, {"bag": "BAG", "best_single_exposure": "Best single exposure", "best_single_exposure_r2": "Single-exposure R²", "arm": "Discovery arm", "triplet": "Triplet", "n_candidates_with_triplet": "Candidates containing triplet", "n_candidates_in_arm": "Candidates in arm", "prevalence_pct": "Prevalence (%)", "omega": "Ω (nats)"})
    shared = publication_frame(read_csv("outputs/dedup/model_comparison/recurrent_triplets_shared.csv"), {"triplet": "Triplet", "structural_synergy_pct": "Synergy arm, structural BAG (%)", "structural_redundancy_pct": "Redundancy arm, structural BAG (%)", "functional_synergy_pct": "Synergy arm, functional BAG (%)", "functional_redundancy_pct": "Redundancy arm, functional BAG (%)", "omega": "Ω (nats)"})
    sheet = TableSheet(workbook, "ST09_RecurrentTriplets", 9, "Recurrent exposome triplets", "Triplets are counted across the top-20 candidates at each of four model levels, with a candidate counted once if it recurs across levels. Prevalence denominators are reported explicitly.")
    sheet.section("Most recurrent triplet by BAG and arm", top)
    sheet.section("Triplets containing the best single exposure", single)
    sheet.section("Triplets shared across both BAGs and arms", shared)
    sheet.finish()


def table_10(workbook: Workbook) -> None:
    subject = read_csv("outputs/dedup/model_comparison/residual_bias/residual_bias_subject_summary.csv")
    subject["bag"] = subject["bag"].str.title(); subject["objective"] = arm_labels(subject["objective"])
    if "diagnosis" in subject:
        subject["diagnosis"] = diagnosis_labels(subject["diagnosis"])
    subject = publication_frame(subject, {c: c.replace("_", " ").title() for c in subject.columns if c not in {"candidate_id", "rung_id"}})
    tests = read_csv("outputs/dedup/model_comparison/residual_bias/residual_bias_diagnosis_tests.csv")
    tests["bag"] = tests["bag"].str.title(); tests["objective"] = arm_labels(tests["objective"])
    for column in ("diagnosis", "contrast"):
        if column in tests:
            tests[column] = diagnosis_labels(tests[column])
    if "groups" in tests:
        tests["groups"] = tests["groups"].str.replace("CN", "HC", regex=False).str.replace("FTD", "FTLD", regex=False)
    tests = publication_frame(tests, {c: c.replace("_", " ").title() for c in tests.columns if c not in {"candidate_id", "rung_id"}})
    paired = read_csv("outputs/dedup/model_comparison/residual_bias/residual_bias_country_paired_diagnosis.csv")
    paired["bag"] = paired["bag"].str.title(); paired["objective"] = arm_labels(paired["objective"])
    paired["contrast"] = diagnosis_labels(paired["contrast"])
    paired = publication_frame(paired, {"bag": "BAG", "objective": "Discovery arm", "outcome": "Outcome", "contrast": "Diagnosis contrast", "n_paired_countries": "Matched countries", "mean_country_difference": "Mean country difference", "median_country_difference": "Median country difference", "ci_low": "95% CI, lower", "ci_high": "95% CI, upper", "permutation_p": "Two-sided sign-flip p", "permutation_exceedances": "Exceedances", "n_draws": "Permutations"})
    sheet = TableSheet(workbook, "ST10_ResidualBias", 10, "Residual bias by diagnosis and country", "Bias is predicted minus observed BAG in years. Equal-country paired contrasts are the primary diagnosis inference; participant-level tests are retained as pooled descriptive sensitivity analyses.")
    sheet.section("Primary within-country diagnosis contrasts", paired)
    sheet.section("Participant-level residual summaries", subject)
    sheet.section("Pooled participant-level sensitivity tests", tests)
    sheet.finish()


def table_11(workbook: Workbook) -> None:
    data = read_csv("outputs/dedup/model_comparison/normative_transfer_results.csv")
    data["bag"] = data["bag"].str.title()
    data["transfer"] = data["transfer"].str.replace("CN", "HC", regex=False).str.replace("FTD", "FTLD", regex=False)
    table = publication_frame(data, {"bag": "BAG", "transfer": "Train → test", "n_scored": "Participants", "r2_baseline": "Baseline R²", "r2_best_syn": "Synergy-arm R²", "r2_best_red": "Redundancy-arm R²", "best_arm": "Best arm", "r2_best": "Best R²", "delta_best_vs_base": "ΔR², best − baseline", "bestVbase_n_countries": "Countries", "bestVbase_median_dR2": "Median country ΔR²", "bestVbase_wins": "Country wins", "bestVbase_wilcoxon_p": "Wilcoxon p", "bestVbase_sign_p": "Sign-test p"})
    sheet = TableSheet(workbook, "ST11_NormativeTransfer", 11, "Normative transfer across diagnostic groups", "Depth-3 models were selected independently within each train-to-test setting. Across-country tests use country as the analysis unit and are two-sided.")
    sheet.section("Diagnostic train-to-test performance", table)
    sheet.finish()


def table_12(workbook: Workbook) -> None:
    estimates = read_csv("outputs/sensitivity/dedup/residual_confounds/lme_estimates_all.csv")
    estimates["bag"] = estimates["bag"].str.title(); estimates["objective"] = arm_labels(estimates["objective"])
    estimates["term"] = estimates["term"].str.replace("FTD", "FTLD", regex=False).str.replace("CN", "HC", regex=False)
    estimates = publication_frame(estimates, {"bag": "BAG", "objective": "Discovery arm", "term": "Fixed effect", "coef": "Estimate", "se": "SE", "ci_lo": "95% CI, lower", "ci_hi": "95% CI, upper", "p_value": "Two-sided Wald p", "n_obs": "Participants", "n_countries": "Countries", "converged": "Converged", "optimizer": "Optimizer"})
    variance = read_csv("outputs/sensitivity/dedup/residual_confounds/lme_variance_all.csv")
    variance["bag"] = variance["bag"].str.title(); variance["objective"] = arm_labels(variance["objective"])
    variance = publication_frame(variance, {"bag": "BAG", "objective": "Discovery arm", "n_obs": "Participants", "n_countries": "Countries", "unadjusted_country_variance": "Unadjusted country variance", "unadjusted_residual_variance": "Unadjusted residual variance", "unadjusted_icc": "Unadjusted ICC", "adjusted_country_variance": "Adjusted country variance", "adjusted_residual_variance": "Adjusted residual variance", "adjusted_icc": "Adjusted ICC", "country_random_intercept_lrt": "Random-intercept LRT", "country_random_intercept_bootstrap_p": "Parametric-bootstrap p", "country_random_intercept_bootstrap_draws": "Bootstrap draws", "adjusted_converged": "Converged"})
    sheet = TableSheet(workbook, "ST12_MixedModels", 12, "Residual-confound linear mixed-effects models", "Maximum-likelihood models of signed participant-level BAG bias include a country random intercept. Continuous predictors are standardized; female and HC are reference levels.")
    sheet.section("Fixed effects", estimates)
    sheet.section("Variance components and country random-intercept tests", variance)
    sheet.finish()


def table_13(workbook: Workbook) -> None:
    domain = read_csv("outputs/sensitivity/dedup/domain_imbalance/manuscript_main_comparison.csv")
    inference = read_csv("outputs/sensitivity/dedup/domain_imbalance/manuscript_matched_inference.csv")
    domain = domain.merge(
        inference[
            [
                "bag",
                "candidate_family",
                "r2_main",
                "r2_alternative",
                "delta_r2_main_minus_alternative",
                "delta_r2_ci_lo",
                "delta_r2_ci_hi",
                "country_cluster_bootstrap_p",
            ]
        ],
        on=["bag", "candidate_family"],
        how="left",
        validate="one_to_one",
    )
    domain["bag"] = domain["bag"].str.title(); domain["score_defined_arm"] = arm_labels(domain["score_defined_arm"])
    domain = publication_frame(domain, {"bag": "BAG", "candidate_family": "Alternative representation", "sensitivity_best_order": "Selected components", "score_defined_arm": "Selected arm", "sensitivity_best_r2": "Alternative R²", "main_best_r2": "Main-analysis R²", "delta_r2_sensitivity_minus_main": "ΔR², alternative − main", "r2_main": "Matched main R²", "r2_alternative": "Matched alternative R²", "delta_r2_main_minus_alternative": "Matched ΔR², main − alternative", "delta_r2_ci_lo": "95% CI, lower", "delta_r2_ci_hi": "95% CI, upper", "country_cluster_bootstrap_p": "Country-cluster bootstrap p"})
    pca_summary = read_csv("outputs/sensitivity/dedup/whole_exposome_pca/manuscript_main_comparison.csv")
    pca_summary["bag"] = pca_summary["bag"].str.title()
    pca_summary = publication_frame(pca_summary, {"bag": "BAG", "best_pc_count": "PCs", "best_pca_r2": "PCA R²", "main_best_r2": "Main-analysis R²", "delta_r2_pca_minus_main": "ΔR², PCA − main", "pc1_variance_fraction": "PC1 variance fraction", "cumulative_variance_fraction": "Cumulative variance fraction", "n_components": "Maximum PCs"})
    incremental = []
    for bag, panel in (("Structural", "a2_struct"), ("Functional", "b2_func")):
        frame = read_csv(f"outputs/figures/dedup/sensitivity/whole_exposome_pca/complete/source_data/whole_exposome_pca_sensitivity_source_data_{panel}_incremental_pcs.csv")
        frame = frame[frame["model_level"] == "d3"].copy(); frame.insert(0, "BAG", bag); incremental.append(frame)
    incremental_frame = publication_frame(pd.concat(incremental, ignore_index=True), {"BAG": "BAG", "pc_n": "PCs", "global_oof_r2": "PCA R²", "level_baseline_r2": "Covariate-baseline R²", "original_best_model_r2": "Main-analysis best R²"})
    sheet = TableSheet(workbook, "ST13_AltRepresentations", 13, "Alternative exposome representations", "Sensitivity analyses replace the original indicators with one representative per domain, within-domain PC1 scores, or 1–10 whole-exposome principal components. Domain-balanced matched comparisons use aligned depth-3 out-of-fold predictions and two-sided country-cluster bootstraps on ΔR² (10,000 draws).")
    sheet.section("Domain-balanced representations", domain)
    sheet.section("Whole-exposome PCA summary", pca_summary)
    sheet.section("Incremental whole-exposome PCA, depth-3", incremental_frame)
    sheet.finish()


def table_14(workbook: Workbook) -> None:
    pool = read_csv("outputs/sensitivity/dedup_neg_o/manuscript_comparison/negative_o_vs_path_only.csv")
    pool["bag"] = pool["bag"].str.title()
    pool["best_objective_path_only"] = arm_labels(pool["best_objective_path_only"])
    pool["best_objective_negative_o"] = arm_labels(pool["best_objective_negative_o"])
    performance = publication_frame(pool, {
        "bag": "BAG",
        "best_syn_r2_path_only": "Path-only synergy R²",
        "best_red_r2_path_only": "Path-only redundancy R²",
        "best_r2_path_only": "Path-only best R²",
        "best_objective_path_only": "Path-only selected arm",
        "best_order_path_only": "Path-only selected size",
        "best_syn_r2_negative_o": "Negative-Ω synergy R²",
        "best_red_r2_negative_o": "Negative-Ω redundancy R²",
        "best_r2_negative_o": "Negative-Ω best R²",
        "best_objective_negative_o": "Negative-Ω selected arm",
        "best_order_negative_o": "Negative-Ω selected size",
    })
    candidate_pool = publication_frame(pool, {
        "bag": "BAG",
        "n_obs_path_only": "Path-only candidates",
        "n_obs_negative_o": "Negative-Ω candidates",
        "n_candidates_removed": "Candidates removed",
        "beta_h_red_path_only": "Path-only redundancy slope",
        "beta_h_syn_path_only": "Path-only synergy slope",
        "beta_interaction_path_only": "Path-only interaction",
        "beta_h_red_negative_o": "Negative-Ω redundancy slope",
        "beta_h_syn_negative_o": "Negative-Ω synergy slope",
        "beta_interaction_negative_o": "Negative-Ω interaction",
        "delta_beta_interaction_negative_o_minus_path_only": "Interaction change",
    })
    summary = read_csv("outputs/sensitivity/dedup_neg_o/manuscript_comparison/negative_o_arm_comparison_summary.csv")
    summary["bag"] = summary["bag"].str.title()
    summary = publication_frame(summary, {
        "bag": "BAG",
        "global_r2_synergy": "Synergy-arm R²",
        "global_r2_redundancy": "Redundancy-arm R²",
        "delta_global_r2_redundancy_minus_synergy": "Redundancy − synergy R²",
        "n_paired_countries": "Matched countries",
        "median_country_delta_r2_redundancy_minus_synergy": "Median country ΔR²",
        "redundancy_country_wins": "Redundancy wins",
        "synergy_country_wins": "Synergy wins",
        "wilcoxon_statistic": "Wilcoxon W",
        "wilcoxon_p_two_sided": "Two-sided p",
        "rank_biserial_redundancy_minus_synergy": "Rank-biserial effect",
    })
    countries = read_csv("outputs/sensitivity/dedup_neg_o/manuscript_comparison/negative_o_arm_comparison_by_country.csv")
    countries["bag"] = countries["bag"].str.title()
    countries = publication_frame(countries, {c: c.replace("_", " ").title() for c in countries.columns if c not in {"syn_candidate_id", "red_candidate_id", "rung_id"}})
    sheet = TableSheet(workbook, "ST14_NegativeOmega", 14, "Negative-Ω arm-definition sensitivity", "The restricted variant additionally requires evaluated Ω<0 for synergy-arm candidates. Candidate-pool summaries are descriptive; the selected arm contrast uses matched-country R² differences and a two-sided paired Wilcoxon signed-rank test.")
    sheet.section("Selected-model performance", performance)
    sheet.section("Candidate pools and diversity slopes", candidate_pool)
    sheet.section("Matched-country arm comparison", summary)
    sheet.section("Country-level matched differences", countries)
    sheet.finish()


def table_15(workbook: Workbook) -> None:
    context = read_csv("outputs/sensitivity/dedup/normative_context/manuscript_transfer_by_context.csv")
    context["bag"] = context["bag"].str.title(); context["objective"] = arm_labels(context["objective"])
    context["condition"] = context["condition"].str.replace("CN", "HC", regex=False).str.replace("FTD", "FTLD", regex=False)
    context = publication_frame(context, {"bag": "BAG", "condition": "Diagnostic context", "model_level": "Selected level", "objective": "Selected arm", "r2": "Selected-model R²", "baseline_r2": "Baseline R²", "delta_r2_vs_baseline": "ΔR²"})
    diversity = read_csv("outputs/sensitivity/dedup/normative_context/manuscript_diversity_by_context.csv")
    diversity["bag"] = diversity["bag"].str.title()
    diversity["condition"] = diversity["condition"].str.replace("CN", "HC", regex=False).str.replace("FTD", "FTLD", regex=False)
    diversity = publication_frame(diversity, {"bag": "BAG", "condition": "Diagnostic context", "n_obs": "Candidates", "beta_h_red": "Redundancy-arm diversity slope", "beta_h_syn": "Synergy-arm diversity slope", "p_h_red": "Redundancy-slope p", "p_h_syn_interaction": "Arm-interaction p", "model_r2": "Model R²"})
    weighting = read_csv("outputs/sensitivity/dedup/diagnosis_balance/manuscript_weighting_comparison.csv")
    weighting["bag"] = weighting["bag"].str.title(); weighting["objective"] = arm_labels(weighting["objective"])
    weighting["diagnosis_label"] = diagnosis_labels(weighting["diagnosis_label"])
    weighting = publication_frame(weighting, {"bag": "BAG", "objective": "Discovery arm", "diagnosis_label": "Diagnosis", "weighted_n": "Participants", "weighted_bias_mean": "Weighted mean bias", "unweighted_bias_mean": "Unweighted mean bias", "delta_bias_mean_weighted_minus_unweighted": "Δ mean bias", "weighted_mae": "Weighted MAE", "unweighted_mae": "Unweighted MAE", "weighted_r2": "Weighted R²", "unweighted_r2": "Unweighted R²"})
    sheet = TableSheet(workbook, "ST15_DiagnosticContext", 15, "Diagnostic context and diagnosis weighting", "Selected-model and diversity results are shown for the pooled and diagnosis-specific contexts. Equal-diagnosis weighting tests whether residual-bias patterns are driven by diagnostic imbalance.")
    sheet.section("Selected models by diagnostic context", context)
    sheet.section("Domain-diversity slopes by diagnostic context", diversity)
    sheet.section("Equal-diagnosis-weighting sensitivity", weighting)
    sheet.finish()


def table_16(workbook: Workbook) -> None:
    folds = read_csv("outputs/sensitivity/dedup/country_region/manuscript_generalization_summary.csv")
    folds["bag"] = folds["bag"].str.title(); folds["geographic_level"] = folds["geographic_level"].str.title()
    folds = publication_frame(folds, {"bag": "BAG", "geographic_level": "Held-out unit", "n_positive_median_r2": "Positive-R² folds", "n_folds": "Folds", "fraction_positive_median_r2": "Positive fraction", "nonpositive_folds": "Non-positive folds", "minimum_median_r2": "Minimum fold R²", "maximum_median_r2": "Maximum fold R²"})
    meta = read_csv("outputs/sensitivity/dedup/country_meta_regression/manuscript_country_meta_regression.csv")
    meta["bag"] = meta["bag"].str.title()
    meta = publication_frame(meta, {"bag": "BAG", "term": "Term", "Coef.": "Estimate", "Std.Err.": "SE", "[0.025": "95% CI, lower", "0.975]": "95% CI, upper", "t": "t", "P>|t|": "Two-sided p", "model_r2": "Model R²", "n_obs": "Countries"})
    sheet = TableSheet(workbook, "ST16_Geography", 16, "Geographical generalisation", "Country-LOCO and region-LOCO summaries use the depth-3 synergy-arm series. Country meta-regressions model country-level out-of-fold R².")
    sheet.section("Country and region folds", folds)
    sheet.section("Country-performance meta-regression", meta)
    sheet.finish()


def table_17(workbook: Workbook) -> None:
    covariates = read_csv("outputs/sensitivity/dedup/education_scanner_baseline/manuscript_covariate_comparison.csv")
    covariates["bag"] = covariates["bag"].str.title()
    covariates = publication_frame(covariates, {"bag": "BAG", "covariate_set": "Covariate specification", "global_oof_r2": "LOCO R²", "reference_r2": "Reference R²", "delta_r2_vs_reference": "ΔR² versus reference", "n_scored": "Participants"})
    residualized = read_csv("outputs/sensitivity/dedup/residualized_bag/manuscript_target_comparison.csv")
    residualized["bag"] = residualized["bag"].str.title()
    residualized = publication_frame(residualized, {"bag": "BAG", "baseline_r2": "Baseline R²", "best_single_r2": "Best single R²", "best_synergy_r2": "Best synergy-arm R²", "best_redundancy_r2": "Best redundancy-arm R²", "delta_single_vs_baseline": "Single ΔR²", "delta_synergy_vs_baseline": "Synergy ΔR²", "delta_redundancy_vs_baseline": "Redundancy ΔR²", "n_candidates": "Candidates"})
    sheet = TableSheet(workbook, "ST17_CovariatesTargets", 17, "Education, scanner and residualized-BAG sensitivity", "All four education/scanner covariate specifications are shown. The residualized-target analysis evaluates BAG after removal of baseline covariate effects.")
    sheet.section("Education and scanner covariates", covariates)
    sheet.section("Residualized BAG target", residualized)
    sheet.finish()


def table_18(workbook: Workbook) -> None:
    caps = read_csv("outputs/sensitivity/dedup/order_cap/order_cap_summary_by_cap_bag_rung.csv")
    caps = caps[(caps["analysis"] == "Pooled") & (caps["rung_id"] == "xgb_tree_d3") & caps["order_cap"].between(5, 30)].copy()
    caps["bag"] = caps["bag"].str.title(); caps["best_objective"] = arm_labels(caps["best_objective"])
    caps = publication_frame(caps, {"bag": "BAG", "order_cap": "Maximum set size", "best_r2": "Best R²", "best_order": "Selected set size", "best_objective": "Selected arm", "best_syn_r2": "Best synergy-arm R²", "best_red_r2": "Best redundancy-arm R²", "syn_minus_red": "Synergy − redundancy R²", "baseline_r2": "Baseline R²", "best_minus_baseline": "Best − baseline R²", "n_candidates": "Candidates"})
    diversity = read_csv("outputs/sensitivity/dedup/order_cap/diversity_r2_betas_by_cap.csv")
    diversity = diversity[(diversity["condition"] == "Pooled") & (diversity["rung_id"] == "xgb_tree_d3") & diversity["order_cap"].between(5, 30)].copy()
    diversity["bag"] = diversity["bag"].str.title()
    diversity = publication_frame(diversity, {"bag": "BAG", "order_cap": "Maximum set size", "n_obs": "Candidates", "beta_h_red": "Redundancy-arm slope", "beta_h_syn": "Synergy-arm slope", "beta_interaction": "Arm interaction", "p_h_syn_interaction_cluster_order": "Set-size-clustered p", "p_h_syn_interaction_perm": "Within-order permutation p", "model_r2": "Model R²"})
    sheet = TableSheet(workbook, "ST18_SetSizeCaps", 18, "Maximum candidate-set-size sensitivity", "Depth-3 pooled results are shown for every maximum set size from 5 to 30. Diversity interaction inference uses set-size-clustered standard errors and within-order label permutations.")
    sheet.section("Selected model by maximum set size", caps)
    sheet.section("Domain-diversity effect by maximum set size", diversity)
    sheet.finish()


BUILDERS = [
    table_01, table_02, table_03, table_04, table_05, table_06,
    table_07, table_08, table_09, table_10, table_11, table_12,
    table_13, table_14, table_15, table_16, table_17, table_18,
]


def validate_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    if len(workbook.sheetnames) != 18:
        raise ValueError(f"Expected 18 worksheets, found {len(workbook.sheetnames)}")
    expected_prefixes = [f"ST{number:02d}_" for number in range(1, 19)]
    for name, prefix in zip(workbook.sheetnames, expected_prefixes):
        if not name.startswith(prefix):
            raise ValueError(f"Worksheet {name!r} does not start with {prefix!r}")
        worksheet = workbook[name]
        expected_title = f"Supplementary Table {int(prefix[2:4])} |"
        if not str(worksheet["A1"].value).startswith(expected_title):
            raise ValueError(f"Unexpected title in {name}: {worksheet['A1'].value!r}")
        if worksheet.max_row < 6 or worksheet.max_column < 2:
            raise ValueError(f"Worksheet {name} contains no substantive table")


def main(output_path: Path = OUTPUT) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for builder in BUILDERS:
        builder(workbook)
    workbook.properties.title = "Supplementary Tables — Exposome higher-order information and brain age gap"
    workbook.properties.subject = "Nature manuscript supplementary tables"
    workbook.properties.creator = "BrainLat exposome–BAG analysis team"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    validate_workbook(output_path)
    print(f"Wrote {output_path} with {len(workbook.sheetnames)} supplementary tables")


if __name__ == "__main__":
    main()
